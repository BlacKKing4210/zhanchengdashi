from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_CELL_TOP_BOTTOM_DXA = 80
TABLE_CELL_START_END_DXA = 120

FONT_BODY = "Microsoft YaHei"
FONT_MONO = "Consolas"
COLOR_HEADING = "2E74B5"
COLOR_HEADING_DARK = "1F4D78"
COLOR_BODY = "202124"
COLOR_MUTED = "5F6B76"
COLOR_TABLE_HEADER = "E8EEF5"
COLOR_CODE_FILL = "F2F4F7"
COLOR_NOTE_FILL = "F7F9FC"


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", TABLE_CELL_TOP_BOTTOM_DXA),
        ("bottom", TABLE_CELL_TOP_BOTTOM_DXA),
        ("start", TABLE_CELL_START_END_DXA),
        ("end", TABLE_CELL_START_END_DXA),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        set_row_cant_split(row)
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def _next_numbering_id(numbering, tag: str, attribute: str) -> int:
    values = []
    for node in numbering.findall(qn(tag)):
        value = node.get(qn(attribute))
        if value is not None and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def create_numbering(document: Document, ordered: bool, start: int = 1) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)

    bullet_markers = ("•", "○", "▪")
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start_node = OxmlElement("w:start")
        start_node.set(qn("w:val"), str(start if level == 0 else 1))
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), f"%{level + 1}." if ordered else bullet_markers[level])
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(540 + level * 360))
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(540 + level * 360))
        indent.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, indent])
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT_BODY)
        fonts.set(qn("w:hAnsi"), FONT_BODY)
        fonts.set(qn("w:eastAsia"), FONT_BODY)
        r_pr.append(fonts)
        lvl.extend([start_node, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
        abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def set_paragraph_numbering(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), str(level))
    num_id_node = num_pr.find(qn("w:numId"))
    if num_id_node is None:
        num_id_node = OxmlElement("w:numId")
        num_pr.append(num_id_node)
    num_id_node.set(qn("w:val"), str(num_id))


def add_hyperlink(paragraph, text: str, target: str):
    relationship_id = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_BODY)
    fonts.set(qn("w:hAnsi"), FONT_BODY)
    fonts.set(qn("w:eastAsia"), FONT_BODY)
    run_properties.extend([fonts, color, underline])
    run_element.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def append_inline(paragraph, value: str) -> None:
    value = value.replace("<br>", "\n").replace("<br/>", "\n")
    cursor = 0
    for match in INLINE_PATTERN.finditer(value):
        if match.start() > cursor:
            run = paragraph.add_run(value[cursor : match.start()])
            set_run_font(run, FONT_BODY)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, FONT_BODY)
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, FONT_MONO, 9.5)
            run.font.color.rgb = RGBColor.from_string(COLOR_HEADING_DARK)
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        set_run_font(run, FONT_BODY)


def set_paragraph_fill(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(COLOR_BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = FONT_BODY
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(COLOR_HEADING_DARK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)

    heading_tokens = {
        "Heading 1": (16, COLOR_HEADING, 18, 10),
        "Heading 2": (13, COLOR_HEADING, 14, 7),
        "Heading 3": (12, COLOR_HEADING_DARK, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in [style.name for style in styles]:
        code_style = styles.add_style("Code Block", 1)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = FONT_MONO
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MONO)
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(COLOR_BODY)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0

    if "Workflow Note" not in [style.name for style in styles]:
        note_style = styles.add_style("Workflow Note", 1)
    else:
        note_style = styles["Workflow Note"]
    note_style.font.name = FONT_BODY
    note_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    note_style.font.size = Pt(10)
    note_style.font.color.rgb = RGBColor.from_string(COLOR_MUTED)
    note_style.paragraph_format.left_indent = Inches(0.18)
    note_style.paragraph_format.right_indent = Inches(0.18)
    note_style.paragraph_format.space_before = Pt(4)
    note_style.paragraph_format.space_after = Pt(6)
    note_style.paragraph_format.line_spacing = 1.2


def configure_document(document: Document, title: str, footer_label: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_paragraph.add_run(footer_label)
    set_run_font(header_run, FONT_BODY, 8)
    header_run.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label_run = footer_paragraph.add_run(f"{footer_label}  |  ")
    set_run_font(label_run, FONT_BODY, 8)
    label_run.font.color.rgb = RGBColor.from_string(COLOR_MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_BODY)
    fonts.set(qn("w:hAnsi"), FONT_BODY)
    fonts.set(qn("w:eastAsia"), FONT_BODY)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    run_properties.extend([fonts, size])
    field_run.append(run_properties)
    field_run.extend([begin, instruction, separate, display, end])
    footer_paragraph._p.append(field_run)

    document.core_properties.title = title
    document.core_properties.author = "Codex Game Studio"
    document.core_properties.subject = "Game development workflow"


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def choose_column_widths(rows: list[list[str]], column_count: int) -> list[int]:
    weights = []
    for column in range(column_count):
        longest = max(
            (len(re.sub(r"[`*_]", "", row[column])) if column < len(row) else 0)
            for row in rows
        )
        weights.append(max(8, min(42, longest)))
    total = sum(weights)
    widths = [max(900, round(CONTENT_WIDTH_DXA * weight / total)) for weight in weights]
    difference = CONTENT_WIDTH_DXA - sum(widths)
    widths[-1] += difference
    if widths[-1] < 900:
        shortage = 900 - widths[-1]
        widths[-1] = 900
        widest = max(range(len(widths) - 1), key=lambda index: widths[index])
        widths[widest] -= shortage
    return widths


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    widths = choose_column_widths(normalized, column_count)
    table = document.add_table(rows=len(normalized), cols=column_count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for row_index, row_data in enumerate(normalized):
        for column_index, value in enumerate(row_data):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            append_inline(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, FONT_BODY, 9)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                set_cell_shading(cell, COLOR_TABLE_HEADER)
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_paragraph(document: Document, value: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    run = paragraph.add_run(value if value else " ")
    set_run_font(run, FONT_MONO, 8.5)
    set_paragraph_fill(paragraph, COLOR_CODE_FILL)


def add_note_paragraph(document: Document, value: str) -> None:
    paragraph = document.add_paragraph(style="Workflow Note")
    append_inline(paragraph, value)
    set_paragraph_fill(paragraph, COLOR_NOTE_FILL)


def markdown_to_docx(source: Path, output: Path, title: str, footer_label: str) -> None:
    document = Document()
    configure_styles(document)
    configure_document(document, title, footer_label)

    lines = source.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    in_code = False
    first_h1 = True
    active_list_type: str | None = None
    active_num_id: int | None = None
    while index < len(lines):
        raw_line = lines[index].rstrip()
        stripped = raw_line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            active_list_type = None
            active_num_id = None
            index += 1
            continue
        if in_code:
            add_code_paragraph(document, raw_line)
            index += 1
            continue

        if stripped.startswith("|"):
            active_list_type = None
            active_num_id = None
            raw_rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = split_table_row(lines[index])
                if not is_table_separator(cells):
                    raw_rows.append(cells)
                index += 1
            add_table(document, raw_rows)
            continue

        if not stripped or stripped == "---":
            active_list_type = None
            active_num_id = None
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            active_list_type = None
            active_num_id = None
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1 and first_h1:
                paragraph = document.add_paragraph(style="Title")
                first_h1 = False
            elif level <= 2:
                paragraph = document.add_paragraph(style="Heading 1")
            elif level == 3:
                paragraph = document.add_paragraph(style="Heading 2")
            else:
                paragraph = document.add_paragraph(style="Heading 3")
            append_inline(paragraph, text)
            set_keep_with_next(paragraph)
            index += 1
            continue

        if stripped.startswith(">"):
            active_list_type = None
            active_num_id = None
            note_parts = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                value = lines[index].strip()[1:].strip()
                if value:
                    note_parts.append(value)
                index += 1
            add_note_paragraph(document, " ".join(note_parts))
            continue

        bullet_match = re.match(r"^(\s*)[-*]\s+(.+)$", raw_line)
        if bullet_match:
            level = min(2, len(bullet_match.group(1).replace("\t", "    ")) // 2)
            if active_list_type != "bullet" or active_num_id is None:
                active_list_type = "bullet"
                active_num_id = create_numbering(document, ordered=False)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            set_paragraph_numbering(paragraph, active_num_id, level)
            append_inline(paragraph, bullet_match.group(2))
            index += 1
            continue

        number_match = re.match(r"^(\s*)(\d+)[.)]\s+(.+)$", raw_line)
        if number_match:
            level = min(2, len(number_match.group(1).replace("\t", "    ")) // 2)
            if active_list_type != "number" or active_num_id is None:
                active_list_type = "number"
                active_num_id = create_numbering(document, ordered=True, start=int(number_match.group(2)))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            set_paragraph_numbering(paragraph, active_num_id, level)
            append_inline(paragraph, number_match.group(3))
            index += 1
            continue

        active_list_type = None
        active_num_id = None
        paragraph = document.add_paragraph()
        append_inline(paragraph, stripped)
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a compact Word workflow guide from Markdown.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--title", required=True)
    parser.add_argument("--footer-label", required=True)
    args = parser.parse_args()
    markdown_to_docx(args.source.resolve(), args.output.resolve(), args.title, args.footer_label)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
