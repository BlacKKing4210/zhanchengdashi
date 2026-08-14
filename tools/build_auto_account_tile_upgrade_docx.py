from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"C:\Users\76398\.codex\skills\game-feature-design-docs\assets\general-feature-design-template.docx"
)
OUTPUT = ROOT / "docs" / "AUTO_ACCOUNT_250_TILE_AND_UPGRADE_RULE_DESIGN_v1.0.docx"

INK = "172033"
BLUE = "2E74B5"
TEAL = "197A78"
GRAY = "5B6574"
LIGHT_BLUE = "EAF3FA"
LIGHT_GOLD = "FFF4D6"
LIGHT_GRAY = "F2F4F7"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Microsoft YaHei"
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 9.3) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    set_run_font(paragraph.add_run(text), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], str(value))
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc: Document, text: str, *, bold: bool = False, color: str = INK) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(text), bold=bold, color=color)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text))


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text))


def add_callout(doc: Document, text: str, color: str = LIGHT_GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), color)
    set_cell_text(table.cell(0, 0), text, bold=True, size=10.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 13, 6),
        ("Heading 2", 12.5, TEAL, 9, 5),
        ("Heading 3", 11, GRAY, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.text = "丛林法则｜自动账号、250 地块与升级生命规则"
    set_run_font(header.runs[0], 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "v1.0｜2026-08-14｜PRODUCER_APPROVED_IMPLEMENTATION_CONTRACT"
    set_run_font(footer.runs[0], 8.5, color=GRAY)


def add_title_page(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(26)
    title.paragraph_format.space_after = Pt(6)
    set_run_font(title.add_run("自动账号、250 地块与升级生命规则"), 23, True, INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run_font(
        subtitle.add_run("账号 ID / 派生密码一键复制｜开局可见 250 金币保底｜固定生命成长"),
        12.2,
        color=BLUE,
    )

    add_table(
        doc,
        ["字段", "内容"],
        [
            ["文档类型", "Game Feature Design General｜Implementation Contract"],
            ["请求编号", "REQ-20260814-AUTO-ACCOUNT-250-HEALTH-RULE"],
            ["功能编号", "GLOBAL｜F-ZC-AUTH-001｜F-ZC-001"],
            ["版本 / 日期", "v1.0 / 2026-08-14"],
            ["批准人与负责人", "制作人：用户｜工程负责人：codex-primary"],
            ["目标平台", "Windows 与 Android 共用 Godot 运行时逻辑"],
            ["规则状态", "制作人明确批准，允许进入实现与验证"],
            ["发布边界", "本地实现；公网加密认证、正式 KDF、阿里云部署和 Android 真机验收另设门禁"],
        ],
        [1.45, 5.1],
    )
    add_callout(
        doc,
        "制作人决策：自动生成账号必须立即拥有可登录的账号 ID 与密码；复制按钮一次复制两者；按钮文案使用“重新登录”；所有势力开局可见 250 金币地块；近战每级生命 +1，远程每两级生命 +1。",
    )
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    for line in (
        "1. 术语、版本与决策",
        "2. 目标、非目标与玩家体验",
        "3. 当前基线与根因",
        "4. 自动账号凭据系统",
        "5. 账号 UI / UE 流程",
        "6. 250 金币地块可见保底",
        "7. 动物升级生命规则",
        "8. 工程映射与数据边界",
        "9. 安全、隐私、日志与异常",
        "10. 美术、音频、埋点与兼容",
        "11. QA 验收矩阵",
        "12. 发布门禁与回滚",
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(line), 10, color=GRAY)
    doc.add_page_break()


def build_document() -> Document:
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    style_document(doc)
    add_title_page(doc)
    add_toc(doc)

    doc.add_heading("1. 术语、版本与决策", level=1)
    add_table(
        doc,
        ["术语", "定义"],
        [
            ["自动账号", "首次安装认证或玩家点击“新建账号”时，由服务器创建且无需玩家先填写名称的账号。"],
            ["账号 ID", "服务器生成、不可修改的 UserID；自动账号的登录账号名与 UserID 完全一致。"],
            ["设备恢复密钥", "客户端首次安装生成的 32 字节随机根凭据，只用于设备恢复和自动账号密码派生，禁止展示或复制。"],
            ["派生密码", "固定域 + 规范化 UserID + 设备恢复密钥经过 SHA-256 得到的 64 位小写十六进制字符串。"],
            ["可见 250 地块", "战斗开始时已经与本方基地相邻、满足解锁条件并显示价格 250 的单位营地。"],
            ["近战 / 远程", "按基础攻击范围分类：≤40 为近战，>40 为远程；技能类型不改变分类。"],
        ],
        [1.45, 5.1],
    )
    add_table(
        doc,
        ["版本", "日期", "变更", "状态"],
        [["v1.0", "2026-08-14", "建立自动凭据、250 可见保底和固定生命成长合同", "制作人已批准"]],
        [0.8, 1.1, 3.75, 0.9],
    )

    doc.add_heading("2. 目标、非目标与玩家体验", level=1)
    doc.add_heading("2.1 主要目标", level=2)
    add_bullet(doc, "首次联网后，玩家不再看到“游客账号（未绑定）/未设置”，而是立即得到账号 ID 和已设置的自动密码。")
    add_bullet(doc, "玩家点击“复制账号密码”后，剪贴板一次得到带标签的账号 ID 与密码，可在另一受支持平台登录同一 UserID。")
    add_bullet(doc, "所有对战地图、所有参战势力开局至少看到一格可解锁的 250 金币单位地块，同时保留初始金矿和 50 金币营地。")
    add_bullet(doc, "动物升级最大生命变成可预测的整数成长，不再叠加动物生命百分比倍率。")
    doc.add_heading("2.2 非目标", level=2)
    add_bullet(doc, "不让服务器保存或返回明文密码、设备恢复密钥、盐或密码哈希。")
    add_bullet(doc, "不修改防御塔按购买次数计价的既有规则，不改变 50/100/250 卡池内容与概率配置。")
    add_bullet(doc, "不改变页面布局、点击矩形、战斗镜头、地图尺寸、金矿定额、群体增益或其他动物技能。")
    add_bullet(doc, "本版本不宣称普通 ENet UDP 已具备正式公网口令传输安全，也不执行阿里云部署或手机出包。")

    doc.add_heading("3. 当前基线与根因", level=1)
    add_table(
        doc,
        ["问题", "当前基线", "根因", "本版处理"],
        [
            ["自动账号无密码", "设备账号记录的 account/salt/password_hash 为空", "安装认证只创建游客档案", "为新档案设置 UserID 登录名与派生密码摘要；旧档案懒迁移"],
            ["复制按钮不可用", "只认本次手动输入的会话密码", "自动登录无法恢复历史口令", "原设备按 UserID 与根密钥重新派生，不读取服务器明文"],
            ["文案不一致", "小按钮显示“重新验证”", "旧安全交互命名遗留", "统一为“重新登录”"],
            ["250 地块缺失", "单位格总体命中率低且无保底", "20% 价格概率只作用于有限单位格", "地图生成后对称/旋转地为每势力补一格开局可见 250 营地"],
            ["升级生命过高", "动物生命先乘等级倍率再叠固定奖励", "旧通用倍率仍作用于生命", "动物生命仅使用基础生命 + 固定奖励；非动物保留旧倍率"],
        ],
        [1.05, 1.55, 1.85, 2.0],
    )

    doc.add_heading("4. 自动账号凭据系统", level=1)
    doc.add_heading("4.1 派生算法", level=2)
    add_callout(
        doc,
        "password = SHA256(" + '"zhanchengdashi:auto-account:v1:" + normalize(UserID) + ":" + recovery_secret' + ")",
        LIGHT_BLUE,
    )
    add_bullet(doc, "normalize(UserID)：去除首尾空白并转为小写；UserID 必须符合 U-<时间戳>-<随机十六进制> 的服务器格式。")
    add_bullet(doc, "输出为 64 位小写十六进制字符串，满足现有 8–72 位密码长度限制。")
    add_bullet(doc, "相同 UserID 与恢复密钥得到稳定密码；不同 UserID 得到不同密码；输出不得等于恢复密钥。")
    add_bullet(doc, "派生实现放在共享纯规则脚本中，但设备恢复密钥只在客户端参与派生；服务器只校验自动账号格式与派生密码摘要。")

    doc.add_heading("4.2 新安装与新建账号", level=2)
    add_number(doc, "客户端以既有 installation_id、refresh_token 与设备恢复流程取得已认证会话；服务端先生成只含 UserID 的游客档案。")
    add_number(doc, "客户端收到 UserID 后，在本机用 recovery_secret 派生自动密码，并通过已认证会话只发送派生密码；该新增操作不发送设备恢复密钥。")
    add_number(doc, "服务器将游客记录原子升级为 account=UserID、auto_generated=true、随机盐和密码摘要；服务端不参与明文派生。")
    add_number(doc, "成功响应只返回 UserID、account、has_password=true、auto_generated=true、auto_password_local=true、会话与资料；不返回派生密码或根密钥。")
    add_number(doc, "客户端仅在本机派生校验成功时按需提供遮罩显示、短时查看与复制；“新建账号”因 UserID 不同而得到不同密码。")

    doc.add_heading("4.3 旧游客档案迁移", level=2)
    add_bullet(doc, "仅当账号记录 account/password_hash 为空且调用者持有该 UserID 的有效会话时允许原子升级；新增 RPC 只接收派生密码，不接收 recovery_secret。")
    add_bullet(doc, "迁移保持 UserID、玩家资料、档案 revision、安装绑定与账号列表不变，只补 account、salt、password_hash、auto_generated。")
    add_bullet(doc, "迁移写盘失败必须原子回滚，旧游客仍可继续使用；不得生成第二份玩家资料。")

    doc.add_heading("4.4 跨平台登录", level=2)
    add_bullet(doc, "玩家在另一 Windows/Android 安装输入复制的账号 ID 与派生密码，沿用现有账号密码登录并绑定新 installation。")
    add_bullet(doc, "新设备不会获得原设备 recovery_secret；登录成功只获得自己的刷新令牌，因此服务器不会泄露原设备根凭据。")
    add_bullet(doc, "命名手动账号继续使用玩家输入密码；自动账号派生逻辑只在 auto_generated=true 时启用。")

    doc.add_heading("5. 账号 UI / UE 流程", level=1)
    add_table(
        doc,
        ["状态", "账号行", "密码行", "小按钮", "复制按钮"],
        [
            ["自动账号（原设备）", "显示 UserID，并标记自动账号", "默认 8 个圆点；10 秒查看派生密码", "查看 / 隐藏", "直接可用，复制账号 ID + 密码"],
            ["命名账号，本次手动登录", "显示账号名", "默认圆点；10 秒查看本次输入", "查看 / 隐藏", "直接可用"],
            ["命名账号，自动登录", "显示账号名", "圆点", "重新登录", "点击后进入重新登录，成功后可复制"],
            ["未登录 / 失效", "显示输入框", "密码输入框", "不显示", "不显示"],
        ],
        [1.35, 1.55, 1.6, 1.0, 1.5],
    )
    add_number(doc, "玩家打开“设置与账号”。")
    add_number(doc, "自动账号状态直接显示账号 ID 与遮罩密码；点击复制后写入“账号ID：…\\n密码：…”。")
    add_number(doc, "剪贴板写入成功后提示 60 秒清理；若玩家在此期间改写剪贴板，客户端不覆盖玩家的新内容。")
    add_number(doc, "应用失焦、进入后台、断线或切换账号时，清除临时显示状态；自动账号密码下一次打开时可从本机重新派生。")
    add_number(doc, "命名账号没有当前会话密码时，小按钮和复制失败路径均进入“重新登录”，不使用“重新验证”文案。")

    doc.add_heading("6. 250 金币地块可见保底", level=1)
    doc.add_heading("6.1 后置条件", level=2)
    add_callout(doc, "每个参战势力在战斗开始时至少有 1 格：site=hall、site_cost=250、无建筑、非 starting_resource、与本方基地相邻且可立即解锁。", LIGHT_BLUE)
    doc.add_heading("6.2 生成顺序", level=2)
    add_number(doc, "按现有地图定义、cell type 权重和 30/50/20 价格规则生成普通地块。")
    add_number(doc, "放置基地、初始金矿、初始 50 营地，并应用每势力额外一座金矿定额。")
    add_number(doc, "检查每个势力基地相邻的非资源地块；若已存在可见 250 营地则不改。")
    add_number(doc, "若不存在，从非基地、非金矿、非 starting_resource 的相邻格中确定性选一格，改为 250 营地。")
    add_number(doc, "1V1/2V2 使用镜像对同时写入；3V3 与自由混战使用六重旋转同时写入，保证价格、站点和可见性公平。")
    add_number(doc, "没有合法候选时地图创建失败并走既有安全回退，禁止静默产出不满足合同的地图。")

    doc.add_heading("6.3 保持不变", level=2)
    add_bullet(doc, "每个势力仍保留一格基地相邻初始金矿、一格 50 金币营地和一格非相邻额外金矿。")
    add_bullet(doc, "`cell_price_pools.csv` 与 `runtime/config/cell_price_pools.json` 的 50/100/250 权重及卡池映射不改。")
    add_bullet(doc, "防御塔静态 site_cost 与购买次数动态价格本次不调整。")

    doc.add_heading("7. 动物升级生命规则", level=1)
    add_table(
        doc,
        ["类别", "判定", "累计生命奖励", "Lv1–8 奖励"],
        [
            ["近战动物", "base_attack_range ≤ 40", "level - 1", "+0, +1, +2, +3, +4, +5, +6, +7"],
            ["远程动物", "base_attack_range > 40", "floor(level / 2)", "+0, +1, +1, +2, +2, +3, +3, +4"],
            ["非动物卡", "building/mine/defense/tower 等标签", "不使用固定奖励", "沿用既有 10.5% 每级倍率"],
        ],
        [1.1, 1.7, 1.6, 2.25],
    )
    add_bullet(doc, "动物最大生命 = round(基础最大生命) + 固定生命奖励；动物生命不再乘等级倍率。")
    add_bullet(doc, "攻击、移动速度、攻击范围和召唤间隔继续使用既有等级倍率。")
    add_bullet(doc, "召唤技能不改变近远程分类；近战召唤动物按近战每级 +1。")
    add_bullet(doc, "卡牌详情、玩家战斗单位、AI 与多人卡组统一调用 `CardRules.card_stats()`，不得另写公式。")

    doc.add_heading("8. 工程映射与数据边界", level=1)
    add_table(
        doc,
        ["文件 / 模块", "职责", "本次变化"],
        [
            ["scripts/shared/account_credential_rules.gd", "纯凭据派生规则", "UserID 校验、自动密码派生、自动账号判断"],
            ["scripts/server/player_account_store.gd", "服务器账号权威存储", "会话内游客账号原子升级、密码摘要校验、auto_generated 状态"],
            ["scripts/network/online_room.gd", "客户端/服务端 RPC 适配", "本机派生；仅发送派生密码；透传 auto_generated 与本地可用状态"],
            ["scripts/app/main.gd", "账号中心 UI 与剪贴板", "自动账号直接查看/复制；“重新登录”文案"],
            ["scripts/app/systems/multiplayer_rules.gd", "全部战斗地图生成", "基地相邻 250 营地镜像/旋转保底"],
            ["scripts/app/systems/card_rules.gd", "卡牌等级统计", "近战/远程固定生命成长"],
            ["config/tables/cell_price_pools.csv", "价格到卡池映射", "只读，确认 250 档存在；本次不改"],
        ],
        [2.45, 1.35, 2.85],
    )

    doc.add_heading("9. 安全、隐私、日志与异常", level=1)
    add_bullet(doc, "服务端账号 JSON、RPC 结果、日志、崩溃输出与管理后台快照均不得包含派生密码、recovery_secret、salt 或 password_hash。")
    add_bullet(doc, "复制内容只写操作系统剪贴板，不写游戏存档、服务器资料或分析事件；埋点只允许记录成功/失败布尔值和错误码。")
    add_bullet(doc, "错误账号与错误密码使用统一 `invalid_credentials`，避免账号枚举；原有会话与安装绑定限制继续生效。")
    add_bullet(doc, "现有普通 ENet RPC 仍会承载口令，正式公网发行前必须迁移到校验证书的加密认证通道。")
    add_bullet(doc, "现有 12,000 轮自制 SHA-256 KDF 仅属开发阶段；正式发行前必须迁移到版本化 Argon2id/PBKDF2/bcrypt 并设置登录限速。")
    add_bullet(doc, "如果自动账号迁移写盘失败，返回 storage_error，保持原记录和安装绑定；不得出现半迁移或重复 UserID。")

    doc.add_heading("10. 美术、音频、埋点与兼容", level=1)
    add_table(
        doc,
        ["领域", "要求"],
        [
            ["UI 美术", "不改布局、按钮尺寸、颜色、字体或点击区域；只替换文案与状态内容。"],
            ["地图美术", "复用既有 hall 图标和 250 数字绘制，不新增资产。"],
            ["音频", "复用 ui_click/ui_confirm/ui_error；不新增音频。"],
            ["埋点", "可选记录 auto_credential_copy、auto_account_migrated、high_price_quota_applied，禁止包含凭据文本。"],
            ["兼容", "Windows 鼠标键盘与 Android 触摸共用现有账号中心；Web/iOS 仍不纳入账号发布声明。"],
            ["可编辑 UE", "本轮保持已批准页面坐标，不新增 UE 页面；若未来重做账号中心布局，再补 Figma/FigJam 页面源。"],
        ],
        [1.3, 5.25],
    )

    doc.add_heading("11. QA 验收矩阵", level=1)
    add_table(
        doc,
        ["编号", "场景", "通过标准"],
        [
            ["AUTH-01", "新 installation 首次认证", "返回 account=UserID、has_password=true、auto_generated=true；复制可用"],
            ["AUTH-02", "派生确定性", "同 UserID 同密钥稳定；不同 UserID 不同；结果不等于根密钥"],
            ["AUTH-03", "跨 installation 登录", "复制的账号 ID + 密码登录到同一 UserID 与资料"],
            ["AUTH-04", "旧游客迁移", "UserID/资料/revision 不变；补齐自动账号凭据；重启后仍可登录"],
            ["AUTH-05", "秘密扫描", "服务器 JSON、响应、日志均搜索不到明文派生密码和 recovery_secret"],
            ["UI-01", "账号中心", "自动账号不显示游客/未设置；小按钮使用查看/隐藏或重新登录"],
            ["UI-02", "复制", "剪贴板精确包含账号ID与密码；60 秒仅清理未被改写的本次内容"],
            ["MAP-01", "1V1/2V2/3V3 全地图，多固定种子", "每势力至少一格基地相邻、可解锁的 250 hall"],
            ["MAP-02", "自由混战，多固定种子", "六势力均有旋转对称的可见 250 hall"],
            ["MAP-03", "地图回归", "基地、初始 50 营地、两座金矿、镜像/旋转、公平性与连通性保持"],
            ["HP-01", "动物 Lv1–8", "近战与远程序列严格匹配本合同"],
            ["HP-02", "非动物与其他属性", "非动物生命和所有非生命属性继续既有倍率"],
            ["ENG-01", "Godot", "缩进检查、全项目解析、专项测试和相关回归全部 0 失败"],
            ["DOC-01", "正式文档", "DOCX 结构检查、页面渲染和 PNG 目检通过"],
        ],
        [0.75, 2.2, 3.6],
    )

    doc.add_heading("12. 发布门禁与回滚", level=1)
    add_callout(
        doc,
        "本地代码与自动化通过只代表 LOCAL IMPLEMENTATION COMPLETE。未通过 TLS 认证通道、生产密码 KDF、阿里云授权部署、Windows/Android 实包跨平台登录和 Android 真机 UI 验收前，不得标记正式发行完成。",
        LIGHT_RED,
    )
    add_bullet(doc, "回滚账号改动时，已迁移的自动账号记录仍是合法命名账号；旧客户端可用账号 ID + 密码手动登录，不得删除玩家数据。")
    add_bullet(doc, "回滚 250 保底时只撤销地图后置条件，不改价格池与卡池数据。")
    add_bullet(doc, "回滚生命规则时必须连同专项测试与正式规则版本一起处理，禁止代码与文档公式不一致。")
    add_bullet(doc, "完成收据必须记录测试命令、结果、提交、推送、未解决发布门禁和工作树保留项。")

    doc.core_properties.title = "自动账号、250 地块与升级生命规则"
    doc.core_properties.subject = "账号凭据复制、地图价格保底与动物升级生命的制作人批准实施合同"
    doc.core_properties.author = "Codex Game Studio"
    doc.core_properties.keywords = "Godot, account, credential, 250, map, health, balance"
    return doc


if __name__ == "__main__":
    document = build_document()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
