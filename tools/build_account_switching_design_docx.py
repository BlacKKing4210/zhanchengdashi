from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ACCOUNT_SWITCHING_DESIGN.docx"

NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(85, 85, 85)
BLACK = RGBColor(32, 37, 43)


def set_font(run, size, color=BLACK, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_text(paragraph, text, size=11, color=BLACK, bold=False):
    run = paragraph.add_run(text)
    set_font(run, size, color, bold)
    return run


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    p_pr.append(borders)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        add_text(paragraph, item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(header, "Jungle Law · Account Switching", 9, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(footer, "v1.0 · 2026-07-18", 9, MUTED)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(4)
add_text(title, "账号切换与新建账号功能设计", 23, NAVY, True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
add_text(subtitle, "服务器权威档案切换、账号摘要与全新游戏进度", 13, MUTED)

metadata = [
    ("状态", "实施中"),
    ("范围", "账号中心、服务器账号存储、联网 RPC"),
    ("数据权威", "专用服务器"),
]
for label, value in metadata:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    add_text(paragraph, f"{label}: ", 11, BLACK, True)
    add_text(paragraph, value)
rule = doc.add_paragraph()
rule.paragraph_format.space_after = Pt(8)
add_bottom_border(rule)

sections = [
    ("1. 目标", [
        "将无效的“注销账号”入口替换为账号切换。",
        "显示当前安装凭据拥有的全部服务器档案，并显示 UserID、段位和动物总数量。",
        "用黄色“新建账号”按钮创建一份全新游戏进度。",
    ]),
    ("2. 服务器边界", [
        "一个安装 ID 最多拥有 8 个独立档案；账号集合只能由有效长期令牌和当前短会话查询。",
        "切换目标必须属于当前安装。服务器撤销旧短会话后签发新短会话，客户端不能伪造账号摘要。",
        "手动账号密码登录成功后，该账号加入当前安装的可切换集合；密码不写入本机凭据或账号列表。",
    ]),
    ("3. 新建账号", [
        "服务器读取 cards 配置生成新手档案：8 个普通动物、新手金矿与防御塔、10 张抽卡券、青铜 1 星。",
        "客户端不提交资源数值；新建成功后立即切换到新档案并完整替换本地进度。",
    ]),
    ("4. 页面与流程", [
        "大厅中央基地 -> 设置与账号 -> 切换账号。",
        "账号列表中当前档案显示“当前”；点击其他档案后同步所选资料并返回账号中心。",
        "列表底部固定显示黄色“新建账号”按钮；连接失败、会话失效、非所属账号或数量上限时保留当前档案并提示原因。",
    ]),
    ("5. 验收", [
        "同一安装能列出全部所属档案，且每项包含 UserID、段位和动物总数量。",
        "新建账号后账号数增加 1，并显示新手资料。",
        "切换后旧档案的卡牌、段位和编组不会混入新档案；切回后原进度保持。",
        "账号存储回归、ENet 联机回归和 Godot 解析通过。",
    ]),
]

for heading, bullets in sections:
    doc.add_heading(heading, level=1)
    add_bullets(doc, bullets)

doc.add_heading("6. Dedicated Server Entry", level=1)
add_bullets(doc, [
    "The Windows Desktop export continues to enter the client scene.",
    "The dedicated_server export enters scenes/server.tscn through a bootstrap scene, so account RPCs run on the UDP 24567 listener.",
])

doc.core_properties.title = "账号切换与新建账号功能设计"
doc.core_properties.subject = "Jungle Law server-authoritative account switching"
doc.core_properties.author = "Codex Game Studio"
doc.save(OUT)
print(OUT)
