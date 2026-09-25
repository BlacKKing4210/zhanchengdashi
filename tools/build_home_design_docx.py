"""Build the editable home implementation contract and explicitly draft UE figures.

Uses the approved general Word template, and the project's existing DOCX helpers.
Online Figma availability and runtime acceptance are independent, recorded gates.
"""
from __future__ import annotations

import hashlib
import argparse
import html
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from build_runtime_gm_panel_docx import (
    TEMPLATE, add_body, add_bullet, add_table, clear_template_body,
    set_run_font, style_document,
)

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/HOMESTEAD_DESIGN_v1.1.docx"
ART = ROOT / "docs/home_ui"
FONT = Path("C:/Windows/Fonts/msyh.ttc")
INK = "253449"
BLUE = "315C84"
GRAY = "5F6B78"


def text_block(draw, xy, value, size=28, fill="#253449", max_width=640):
    font = ImageFont.truetype(str(FONT), size)
    x, y = xy
    for paragraph in value.split("\n"):
        line = ""
        for char in paragraph:
            if line and draw.textlength(line + char, font=font) > max_width:
                draw.text((x, y), line, font=font, fill=fill)
                y += size + 10
                line = ""
            line += char
        draw.text((x, y), line, font=font, fill=fill)
        y += size + 10
    return y


def svg_text(draw, xy, value, size=28, fill="#253449", max_width=640):
    """Use the same line wrapping as the PNG draft; all text stays editable."""
    font = ImageFont.truetype(str(FONT), size)
    x, y = xy
    result = []
    for paragraph in value.split("\n"):
        lines = []
        line = ""
        for char in paragraph:
            if line and draw.textlength(line + char, font=font) > max_width:
                lines.append(line)
                line = ""
            line += char
        lines.append(line)
        for line in lines:
            result.append(f'<text x="{x}" y="{y+size}" font-family="Microsoft YaHei" font-size="{size}" fill="{fill}">{html.escape(line)}</text>')
            y += size + 10
    return result


def build_figures(art_dir=ART):
    """Local SVG source and matching PNG are review drafts, never Figma proof."""
    art_dir.mkdir(parents=True, exist_ok=True)
    pages = [
        ("H00 大厅入口", "顶部  金币拥有量    抽卡券拥有量", "保留当前大厅主体", "左下角 家园  [红点]", "仅可领收益亮红点；领取成功立即消失"),
        ("H01 家园主页", "顶部  金币拥有量    抽卡券拥有量", "彩色六边形建筑与边线道路\n中心主城堡  居民随机活动", "拖动地图；有收益时 [领取收益]", "有收益自动进入 H04；锁格见 H02；已有格查看 H03"),
        ("H02 锁格点击反馈", "邻接格  类型剪影  金币价格", "拖动时连续显示，边缘由 UI 遮挡\n非邻接格只显示中性六边形", "价格不可读先居中，不扣款", "可读再点建造；非邻接不聚焦；不打开详情"),
        ("H03 已拥有详情", "温馨小屋    每日产出", "品质底色道具格\n正式图标与数量  明日开始产出", "选择其他建筑时更新", "无容量活动说明或重复道具名称；城堡两项收益"),
        ("H04 收益清单", "家园收益    最近 1～3 天", "品质底色  道具图标与数量\n城堡  金币 ×300   抽卡券 ×3", "[上一页] 1/2 [下一页]   [全部兑换]", "同道具同品质合并；成功清红点；关闭仍可领"),
        ("H05 兑换结果", "家园收获已到账", "金币图标  合计随机金币\n抽卡券图标  合计获得数量", "[收下]", "回 H01 更新余额与红点；同日重进不重复发奖"),
        ("H06 请求反馈", "保留原家园或收益面板", "防重复提交；操作失败短提示\n不显示常驻连接条或重试连接按钮", "操作完成或短提示消失", "失败不扣不发，重新操作不重复奖励"),
        ("H07 休息与重进", "夜晚  动物沿道路回家", "抵达门口后隐藏\n完成活动冒短暂表情气泡", "[大厅]    关闭与后台不累计活动奖励", "再进 H01 读取当前日；动物表现不控制结算"),
    ]
    width, height = 1760, 2500
    img = Image.new("RGB", (width, height), "#f8fafb")
    d = ImageDraw.Draw(img)
    text_block(d, (56, 28), "家园 UE 总流程  v1.1", 42, max_width=1640)
    text_block(d, (56, 88), "低保真评审草稿  本地 SVG 可编辑  Figma 保存与回读待完成", 26, fill="#76613f", max_width=1640)
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="#f8fafb"/>']
    svg.extend(svg_text(d,(56,28),"家园 UE 总流程  v1.1",42,max_width=1640))
    svg.extend(svg_text(d,(56,88),"低保真评审草稿  本地 SVG 可编辑  Figma 保存与回读待完成",26,fill="#76613f",max_width=1640))
    for index, (title, header, body, controls, note) in enumerate(pages):
        x = 56 + (index % 2) * 850
        y = 164 + (index // 2) * 580
        w, h = 798, 516
        d.rounded_rectangle((x, y, x+w, y+h), 22, fill="#ffffff", outline="#b6c5d1", width=3)
        d.rounded_rectangle((x+14, y+14, x+w-14, y+78), 12, fill="#dfeaf0")
        text_block(d, (x+32, y+25), title, 32, max_width=w-64)
        text_block(d, (x+32, y+106), header, 27, max_width=w-64)
        d.rounded_rectangle((x+30, y+164, x+w-30, y+308), 12, fill="#edf1e3", outline="#cbd5b9", width=2)
        text_block(d, (x+50, y+182), body, 29, max_width=w-100)
        text_block(d, (x+32, y+332), controls, 28, fill="#315c84", max_width=w-64)
        text_block(d, (x+32, y+404), note, 25, fill="#5f6b78", max_width=w-64)
        svg.append(f'<g id="{title.split()[0]}"><rect x="{x}" y="{y}" width="{w}" height="{h}" rx="22" fill="white" stroke="#b6c5d1" stroke-width="3"/>')
        svg.append(f'<rect x="{x+14}" y="{y+14}" width="{w-28}" height="64" rx="12" fill="#dfeaf0"/>')
        svg.append(f'<rect x="{x+30}" y="{y+164}" width="{w-60}" height="144" rx="12" fill="#edf1e3" stroke="#cbd5b9" stroke-width="2"/>')
        for offset, txt, size, color in [(25,title,32,"#253449"),(106,header,27,"#253449"),(182,body,29,"#253449"),(332,controls,28,"#315c84"),(404,note,25,"#5f6b78")]:
            svg.extend(svg_text(d,(x+(50 if txt==body else 32),y+offset),txt,size,fill=color,max_width=w-(100 if txt==body else 64)))
        svg.append('</g>')
    # Exact transition metadata supplements the visible action/return notes on each page.
    links = [(0,1,"点击家园"),(1,2,"点未解锁地块"),(1,3,"点已拥有地块"),(1,4,"有收益 / 点收益"),(4,6,"确认兑换"),(6,5,"提交成功"),(5,1,"收下并返回"),(6,1,"取消 / 返回"),(7,1,"次日或再次进入")]
    svg.append('<metadata id="transition-map">'+html.escape(json.dumps(links,ensure_ascii=False))+'</metadata></svg>')
    img.save(art_dir / "home_ue_flow_draft.png")
    (art_dir / "home_ue_flow_draft.svg").write_text('\n'.join(svg), encoding="utf-8")
    # Split the dense overview for readable full-page Word placement.
    img.crop((0, 0, width, 1290)).save(art_dir / "home_ue_pages_a.png")
    img.crop((0, 1275, width, 2500)).save(art_dir / "home_ue_pages_b.png")
    register = {
        "feature_id":"F-ZC-HOME-001", "source_version":"1.1", "owner":"codex-primary",
        "formal_tool":"Figma/FigJam per project AGENTS.md", "review_state":"Pending",
        "reason":"Figma whoami tool call failed with MCP transport HTTP request error on 2026-09-24; no authenticated edit or save/readback evidence.",
        "figma_url":None, "file_id":None, "page_or_board":None, "object_refs":[],
        "coverage":[p[0] for p in pages], "local_source_backup":"docs/home_ui/home_ue_flow_draft.svg",
        "preview_png":["docs/home_ui/home_ue_pages_a.png","docs/home_ui/home_ue_pages_b.png"],
        "warning":"Local editable review draft, not a passed Figma deliverable. Runtime acceptance is recorded separately."
    }
    (art_dir / "artifact_register.json").write_text(json.dumps(register,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def heading(doc, title):
    doc.add_heading(title, level=1)


def page(doc, title):
    doc.add_page_break()
    heading(doc, title)


def table(doc, heads, rows, widths=None):
    return add_table(doc, heads, rows, widths)


def body(doc, text):
    add_body(doc, text)


def build_document(art_dir=ART):
    build_figures(art_dir)
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    style_document(doc)
    for toc_name in ("toc 1", "toc 2", "toc 3"):
        if toc_name in doc.styles:
            toc_style=doc.styles[toc_name]
            toc_style.paragraph_format.space_before=Pt(0)
            toc_style.paragraph_format.space_after=Pt(2)
            toc_style.paragraph_format.line_spacing=1.0
            toc_style.paragraph_format.keep_with_next=False
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "丛林法则  家园模拟经营  F-ZC-HOME-001"
    set_run_font(sec.header.paragraphs[0].runs[0], 8.5, color=GRAY)
    foot=sec.footer.paragraphs[0]
    foot.text="v1.1  2026-09-25  第 "
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); foot._p.append(field)
    foot.add_run(" 页")
    for run in foot.runs: set_run_font(run,8.5,color=GRAY)
    title_style=doc.styles["Title"]
    title_style.font.name="Microsoft YaHei"; title_style.font.size=Pt(24)
    title_style.font.color.rgb=RGBColor(0,0,0)
    title_style.paragraph_format.space_before=Pt(0)
    title_style.paragraph_format.space_after=Pt(12)
    for element in list(title_style._element.xpath('.//w:pBdr')):
        element.getparent().remove(element)
    doc.add_paragraph("家园模拟经营玩法设计",style="Title")
    body(doc,"玩家从大厅左下角进入家园，解锁彩色六边形建筑，观察动物沿道路生活，并领取最近最多三天的建筑收益。本稿把已批准规则落实到程序、数值、界面和测试；功能是否完成以运行时证据为准。")
    table(doc,["版本控制","内容"],[
        ["文件状态","实现合同 v1.1 可读性修订；运行验收以本次回执为准；Figma 正式图稿 Pending；未做 Android 真机验收"],
        ["编号与分类","REQ-20260925-HOME-GACHA-READABILITY；F-ZC-HOME-001 / F-ZC-GACHA-NEW-001；IMPLEMENTATION_CONTRACT / MATERIAL；通用模板"],
        ["负责人","制作人：用户；实施责任：codex-primary；文档编制：Codex"],
        ["创建与完成","创建 2026-09-24，更新 2026-09-25；本稿定义修订规则，历史验收不作为本次通过结论"],
        ["正式基线","Godot 4.6.2；720 × 1280 竖屏；现有 B 美术皮肤、战斗六边形、城堡和动物"],
        ["实施范围","家园与抽卡界面可读性、拖动稳定显示、领取红点及翻牌演出；本次不改经济、服务端或存档，不部署或导出设备包"],
    ],[1.15,5.5])
    table(doc,["版本","日期","更新内容"],[["1.0","2026-09-24","新增家园、最多补领3天、解锁赠券、新地次日收益及新动物演出。"],["1.1","2026-09-25","既有修订：固定圈数外观、直接建造、建筑放大与柔和音效。最新修订：拖动不闪、仅邻接格显示信息、日产品质格、领奖清红点、原比例结果卡与翻牌；取代可买地块红点及旧品质文字规则。"]],[0.6,1.0,5.05])
    table(doc,["计划责任","开始","预计与关闭条件"],[
        ["设计与文档","2026-09-24","本次任务；规则、配置与页面一致，Word 全页可读"],
        ["美术与实现","2026-09-24","本次任务；正式图标、真实入口、各行为可运行"],
        ["QA 与制作人","实现完成后","作用域测试和玩家截图通过；制作人复核另记"],
    ],[1.3,1.1,4.25])

    doc.add_page_break()
    toc_title=doc.add_paragraph("目录")
    set_run_font(toc_title.runs[0],16,True,BLUE)
    p=doc.add_paragraph()
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    code=OxmlElement("w:instrText"); code.set(qn("xml:space"),"preserve"); code.text=' TOC \\o "1-3" \\h \\z \\u '
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    for e in (begin,code,sep): p.add_run()._r.append(e)
    p.add_run("自动目录将在 Word 更新字段后显示。")
    p.add_run()._r.append(end)
    settings=doc.settings._element
    upd=OxmlElement("w:updateFields"); upd.set(qn("w:val"),"true"); settings.append(upd)

    page(doc,"1 术语与设计目的")
    table(doc,["术语","定义"],[
        ["真实日","按北京时间 UTC+8 的自然日；在线以服务端时钟为准"],
        ["视觉日夜","240 秒循环的场景表现；与真实日收益无关"],
        ["预估价值 V","内部兑换参数，玩家界面不展示；初圈 5～15，后续上限 100"],
        ["已拥有与可解锁","拥有状态已持久化；可解锁须邻接已拥有地块且金币足够"],
        ["格式语义","红色为强约束；蓝色为本次新增；绿色为配置；状态同时以文字表达"],
    ],[1.3,5.35])
    doc.add_heading("1.1 主要与次要目标",level=2)
    for s in ["让玩家用现有金币扩展一个可观察、可成长的动物家园。","邻接已拥有地块的锁格通过剪影和价格表达可扩建信息；解锁即时得到一张抽卡券。","动物有昼夜生活节奏；行走、进出建筑、正面心情占多数均能直接看见。","每日登录有稳定城堡基础收益，最多三天补领降低漏登损失。"]: add_bullet(doc,s)
    doc.add_heading("1.2 非本期范围",level=2)
    body(doc,"本期不做建筑升级、搬迁、拆除退款、生产队列、访客社交、居民饥饿惩罚或离线活动模拟；不改变战斗经济。服务器部署、远程数据迁移、Android 发行需要单独的已批准环境与验收。")
    heading(doc,"2 功能概述与系统框架")
    table(doc,["层次","玩家结果与规则归属"],[
        ["入口与浏览","左下角家园按钮可进入；顶部显示现有金币和抽卡券；地图可拖动；保留手势缩放，不显示镜头控制按钮。"],
        ["扩建","邻接锁格显示类型剪影与价格，非邻接只显示中性六边形；价格可读时点击校验后解锁、券 +1。邻接格价格不可读先居中，不扣款。"],
        ["生活","已拥有动物种类成为居民 → 分房 → 沿边线道路活动 → 到门口隐藏或在建筑附近停留 → 表情。"],
        ["领取","读取真实日与家园 → 展示最多 3 天产物 → 点击全部兑换 → 一次结算 → 展示总金币和券。"],
        ["权威与恢复","在线账号由服务端保存状态与经济；游客本地隔离。任何 UI 刷新、重进或重试不能重复领奖。"],
    ],[1.1,5.55])

    page(doc,"3 UE 总流程与页面前半")
    body(doc,"页面 ID 是实现和 QA 的共同索引。以下为低保真页面示意，不是新增美术方向；正式 Figma 页面保存与对象回读为 Pending，原因见第 12 章。")
    doc.add_picture(str(art_dir/"home_ue_pages_a.png"),width=Inches(6.55))
    table(doc,["起点","动作或条件","去向与反馈"],[
        ["H00 大厅","点击家园","H01；先读取权威快照，成功后决定是否自动打开 H04"],
        ["H01 家园","点锁地 / 已拥有地","H02锁格点击反馈 / H03已拥有详情；拖动不误触购买"],
        ["H02 锁格点击","邻接且价格可读时建造","检查通过回 H01，券 +1；邻接格价格不可读先居中；非邻接不聚焦、不打开详情"],
        ["H03 已有详情","选择其他已有格","更新详情；H02无详情面板，不保留隐形按钮"],
    ],[1.25,1.65,3.75])
    body(doc,"拖动时地块、建筑与邻接锁格信息连续绘制，仅由地图上下的界面遮罩遮挡，不按地块中心是否越过边缘整组隐藏。邻接锁格价格被遮挡或缩小到不可读时，首击只居中显示价格，不扣款；次击建造。非邻接格不显示类型、价格或建筑预览，点击也不进入待购买聚焦。")

    page(doc,"4 UE 总流程与页面后半")
    doc.add_picture(str(art_dir/"home_ue_pages_b.png"),width=Inches(6.55))
    table(doc,["起点","动作或条件","去向与反馈"],[
        ["H01","新日有收益 / 点收益","H04，列出物品、天数、城堡金币和券；无收益时不显示领取按钮"],
        ["H04","全部兑换 / 关闭","提交进入 H06；关闭回 H01，收益仍可领取"],
        ["H06","成功 / 失败 / 重试","成功 H05；失败保留可读原因。重试先查询交易或快照，不能重抽奖励"],
        ["H05","收下 / 中断再进","回 H01；读到已到账余额。结果面板不会二次发奖"],
        ["H07","夜晚 / 关闭 / 重进","夜晚安排娱乐与休息；退出仅停止表现；重进读取现状并判断真实日"],
    ],[1.15,1.85,3.65])

    page(doc,"5 参考与准确数据来源")
    body(doc,"来源优先级：用户当前请求与补充 → 本次实施来源 → 权威 CSV 与导出 JSON → 本文规则及页面说明。参考截图只用于理解六边形建筑、地块价格、动物出入与气泡，不复制截图资产、地图或商业识别。无动态参考视频。")
    table(doc,["来源","定位与用途"],[
        ["需求来源","docs/HOME_GACHA_POLISH_IMPLEMENTATION.md 的 9月25日可读性修订；docs/receipts/REQ-20260925-HOME-GACHA-READABILITY.md；用户最新文字与截图"],
        ["经济表","config/tables/home_economy.csv；运行导出 runtime/config/home_economy.json"],
        ["建筑表","config/tables/home_buildings.csv；运行导出 runtime/config/home_buildings.json"],
        ["规则与状态","scripts/shared/home_rules.gd；定价、目录、状态归一化、邻接、快照、结算"],
        ["应用与表现","scripts/app/main.gd；scripts/app/ui/home_view.gd；scripts/app/systems/home_simulation.gd"],
        ["在线权威","scripts/server/player_account_store.gd；scripts/server/player_account_profile_adapter.gd；scripts/network/home_rpc.gd"],
        ["配置验证","config/schema/config_schema.json；tools/validate_config.py；tools/export_config.py；CSV 改动后立即验证并导出"],
    ],[1.1,5.55])
    table(doc,["表与字段","默认值和含义"],[
        ["home_economy::max_rings / first_ring_base / ring_multiplier","8 / 100 / 3；最多 8 圈，中心不计入"],
        ["price_min_pct / price_max_pct","0.5 / 1.5；同圈整数金币范围"],
        ["catchup_days / timezone_hours","3 / 8；最近三天，含今天"],
        ["castle_gold / castle_tickets / unlock_tickets","100 / 1 / 1；每日城堡金币、每日城堡券、每地解锁券"],
        ["linear_limit / linear_divisor / log_base_value / log_growth / max_daily_value","150 / 10 / 15 / 12 / 100；内部价值曲线"],
        ["coin_min_pct / coin_max_pct / ticket_value / ticket_value_share","0.8 / 1.2 / 100 / 0.1；随机金币和额外券概率"],
        ["day_cycle_seconds / positive_mood_chance","240 / 0.85；视觉昼夜与正面气泡比例"],
        ["home_buildings 字段","id / name / item_id / item_name / color / capacity_min / capacity_max / active_period / activity"],
    ],[3.3,3.35])

    body(doc,"外观配置：config/tables/home_visuals.csv → runtime/config/home_visuals.json。id/building_type/ring/ art_path 为固定圈图映射；width/height 为可见绘制尺寸；crop_* 为归一化有效透明边界。第1圈138×136起，第8圈166×157；原图不改像素，按有效范围等比显示。")

    page(doc,"6 地块定价与建筑日产")
    body(doc,"坐标采用轴向 q、r。圈数 R = max(|q|, |r|, |q+r|)，中心 R=0 是默认城堡。圈数只决定固定建筑外观，不存在建筑升级。类型、价格、容量为固定目录，重启、重进和刷新不改变；首圈按六个方向分配并覆盖四类建筑。只解锁邻接已拥有地块；展示最远已拥有圈的下一圈，上限 8 圈，共 217 格含城堡。")
    body(doc,"R≥1 时基准价 B = 100 × 3^(R−1)，整数价格 C ∈ [round(0.5B), round(1.5B)]。例如第一圈 50～150，第二圈 150～450，第三圈 450～1350；圈间端点可相同。")
    table(doc,["建筑","辨识与活动","每日产出"],[
        ["主城堡","米黄色，默认拥有；容纳所有无床位居民","固定 100 金币 + 1 抽卡券"],
        ["居住 温馨小屋","蓝色 82add3；夜间休息；容量稳定为 1～5","软绒枕头 soft_pillow ×1"],
        ["吃饭 森林餐厅","橙色 dea570；白天随机进餐","美味餐篮 meal_basket ×1"],
        ["娱乐 星光剧场","紫色 b8a0d0；夜间随机娱乐","音乐盒 music_box ×1"],
        ["运动 活力球场","绿色 a9be74；白天随机运动","运动奖章 sports_medal ×1"],
    ],[1.35,2.65,2.65])
    doc.add_heading("6.1 内部价值与兑换",level=2)
    body(doc,"每个普通建筑每天产 1 件对应道具，不用进出次数决定数量。内部价值：C≤150 时 V=round(C/10)；C>150 时 V=round(15+12×ln(C/150))；最终夹取 1～100。第一圈 V=5～15，后续比例随价格升高而降低，价值非递减且封顶 100。")
    table(doc,["价格示例 C","内部 V","每件金币随机区间","每件额外 1 券概率"],[
        ["50","5","4～6","0.5%"],["100","10","8～12","1%"],["150","15","12～18","1.5%"],["450","28","23～33","2.8%"],["1350","41","33～49","4.1%"],["达到价值上限","100","80～120","10%"],
    ],[1.4,0.85,2.3,2.1])
    body(doc,"每件道具分别抽取整数金币 G∈[ceil(0.8V), floor(1.2V)]，并独立判定一张额外券，概率 P=V×10%÷100=V÷1000。金币与券可以同时获得。合并显示同类道具不改变逐件抽取；每次成功结算只使用并保存一次随机结果，UI 不显示 V 或兑换概率。")

    page(doc,"7 真实日补领与结算事务")
    body(doc,"真实日 D=floor((服务端 Unix 秒数+8×3600)/86400)。首次初始化家园在当天可领城堡收益；普通建筑只在 D>unlock_day 的日子产出。可结算日从 max(started_day, last_claim_day+1, D−2) 到 D；最多 3 天，超过窗口的漏登日过期，不事后补回。")
    table(doc,["例子","结果"],[
        ["9 月 24 日第一次进入","可领当天城堡 100 金币和 1 券；当天新开地块立刻 +1 券，日产从 25 日开始"],
        ["24 日已领，27 日再进入","补领 25、26、27 日，共 3 天；各地块只计其已生效日"],
        ["24 日已领，30 日再进入","只领 28、29、30 日；25～27 日过期"],
        ["26 日解锁，27 日领取 25～27 日","该地只产 27 日的 1 件；不补产解锁前或解锁当天"],
        ["同日反复进出或重复确认","领取游标未变化则拒绝第二次结算；不重新随机、不重复加资源"],
    ],[2.25,4.4])
    doc.add_heading("7.1 解锁事务",level=2)
    body(doc,"解锁请求携带会话、地块 ID 与档案版本，不由客户端提交价格或奖励。服务端校验账号、版本、合法地块、未拥有、邻接与余额，一次保存 wallet_gold−C、owned[plot_id]=D、gacha_tickets+1 并推进版本。重复返回已拥有，不重复扣款或赠券；保存失败一致回滚。")
    doc.add_heading("7.2 领取事务与不重抽",level=2)
    body(doc,"先展示每日快照，再确认全部兑换。请求携带档案版本与快照日；服务端用自己的当前日校验，过期快照返回最新状态供重试。通过后逐件抽取、加城堡收益，原子保存余额、last_claim_day、last_reward 与版本；成功才回报到账。重放已完成请求返回同一回执，不重抽或二次发奖。")
    table(doc,["字段","含义与恢复"],[
        ["home.version / started_day","家园数据版本 / 首次初始化日"],
        ["home.owned[plot_id]","已拥有地块与解锁日；必须含 0,0 主城堡"],
        ["home.last_claim_day / last_reward","结算游标 / 已结算日、物品、金币和券回执"],
        ["wallet_gold / gacha_tickets","沿用账号余额和正式图标；与家园写入同一次提交"],
        ["profile_revision","用于家园请求冲突检查；慢回包合并期间本地新增金币和券，完成回调不能重复加减"],
        ["在线与游客","在线以账号独立权威档案和服务端时间结算；游客仅独立本地存档。切换账号清空旧 UI/在途回应，不合并状态"],
    ],[2.45,4.2])
    body(doc,"跨日停留刷新可领状态；跨日或版本变化时重新查询快照，不静默结算旧面板。在线时钟由服务端控制，已领日不能补造；游客仍受本机时钟约束。视觉昼夜不会产出或消耗经济资源。")

    page(doc,"8 居民日夜与道路行为")
    body(doc,"居民来自账号已拥有的动物种类，每个已解锁动物卡种 1 位居民，不因重复卡牌数量增殖。首次没有居住建筑时全部住主城堡。按稳定顺序分配住宅床位，满额后剩余居民回城堡；城堡容量不设有限床位。")
    table(doc,["状态","触发与动作","结束与下一步"],[
        ["室内休息","夜晚安排入住所属住宅或城堡；室内隐藏","白天从其门口出现，进入空闲"],
        ["空闲等待","随机等待，避免全体同一帧行动","白天选餐厅/运动；夜晚选娱乐后休息"],
        ["沿路移动","起点门口接入六边形顶点，沿共享边路网寻路","到目的门口，进入室内或建筑周边活动"],
        ["活动中","白天随机吃饭、运动；夜间随机娱乐","到时完成，产生心情气泡，再选择下一状态"],
        ["心情反馈","完成行为时抽取；正面占 85%","气泡在头顶短暂出现后淡出，不遮挡关键价格"],
    ],[1.15,2.8,2.7])
    doc.add_heading("8.1 路网几何合同",level=2)
    body(doc,"六边形边线就是路。相同世界位置的顶点去重，每条边连接两个顶点，居民只能沿这些边连续行走；建筑出入口与最近路点之间允许一小段过门路线。禁止从一个建筑中心直线穿过其他地块。进入建筑后完全隐藏；在建筑周围的行为仍处于可通行边缘。")
    doc.add_heading("8.2 缺建筑与中断",level=2)
    table(doc,["条件","处理"],[
        ["缺餐厅、娱乐或运动建筑","略过该类活动，改为空闲或回住处；不报错、不扣收益、不制造永久等待"],
        ["住宅容量不足","稳定分配现有床位，超员居民住城堡"],
        ["路径不可达或目录无效","取消该次目标，回有效道路/住处；不瞬移穿地块"],
        ["昼夜切换时仍在移动或活动","完成当前安全过渡后更新目标；夜晚最终能休息，白天能出门"],
        ["离开家园、切账号、应用后台","停止场景模拟，清理气泡和路径；重新进入按当前居民及建筑重建，不结算离线行为"],
        ["大量居民","角色辨认、气泡数量和刷新负载须受控；优化不能改变拥有动物数据或经济产量"],
    ],[2.05,4.6])

    page(doc,"9 UI 元素与特殊状态")
    body(doc,"主目标 720×1280，同时检查 360×640。720 基准布局：顶部资源区 y=0～190，地图 y=190～950，底部详情 y=962～1128，导航保持既有位置。复用现有皮肤和正式图标；邻接锁格显示剪影与价格，非邻接锁格只有中性六边形；已有建筑是地块主体。")
    table(doc,["元素","显示与交互","来源与异常"],[
        ["家园按钮与红点","只提示可领取的每日收益；领取成功立即消失","待领取快照为依据；可买地块不触发；请求中或失败仍可领则保留，次日可领时恢复"],
        ["顶部资源","金币与抽卡券图标及余额；必要数量换行或留足空间","现有账号资源；服务不可用不展示旧账号数值为当前"],
        ["锁定地块","只在邻接时显示类型剪影和价格；拖动不整组显隐","邻接格价格不可读先居中、不扣款；非邻接不显示信息、不聚焦；拖动不购买"],
        ["已拥有详情","建筑身份、每日产出标题、下方道具格；必要时短示明日开始产出","每格正式图标与数量，品质用底色；去除容量/活动及重复名称、日率文字；城堡列两项收益"],
        ["收益清单","覆盖天数、道具品质格、城堡收益、翻页、全部兑换","同道具同品质合并、不同品质分格；每页最多4类；数量与逐件结算保持一致"],
        ["兑换结果","实际金币与券的总数；收下只关闭","已持久化 last_reward；重复打开不调用发奖"],
        ["加载与错误","请求时禁用重复点击；网络/保存失败可读并可重试","RPC / 存储结果；不使用乐观扣款或成功提示掩盖失败"],
        ["地图操作","地图与信息在界面遮罩后连续移动，无边缘闪现","保留拖动和手势；无缩放/回中按钮及连接横条；双指不误买，弹窗拦截输入"],
    ],[1.2,2.85,2.6])
    body(doc,"道具格按既有内部预估日兑换价值 V 着色：V<15 为绿，15≤V<40 为蓝，40≤V<70 为紫，V≥70 为金。只改变展示，不改变价值、奖励、容量和活动规则；格内不显示品质名称或内部价值。")
    doc.add_heading("9.1 必须覆盖的页面状态",level=2)
    body(doc,"覆盖地图各边缘鼠标/触摸拖动、解锁后的邻接更新、金币不足/禁用；详情的次日生效、数量与品质；收益1天/3天、最大列表、成功清红点、失败保留、次日恢复；连接失败重查、切账号与重进。")
    body(doc,"最大值不得用縮小文字、裁切或省略号解决。优先扩大容器、换行、减列与分页；操作按钮不得被长清单挤出安全区域。说明文案和错误必须保留可读短句；新图标不能使用 Emoji 或字体符号冒充正式资源。")

    page(doc,"10 美术音频埋点与关联")
    table(doc,["类别","要求与交付"],[
        ["沿用美术","当前已解锁战斗六边形、动物、正式金币和券图标；新家园城堡沿用身份并独立生成。"],
        ["新增道具图标","软绒枕头、美味餐篮、音乐盒、运动奖章；透明背景、稳定资源 ID；与当前风格一致，在实际 UI 尺寸可区分。"],
        ["建筑表现","四类各生成 1～8 圈独立外观，共 32 张，另生成主城堡。第1圈最简单，越外圈体块越大；无升级玩法。"],
        ["心情气泡","头顶象牙白对话气泡，短尖角指向动物；气泡内只放一个 Emoji：😊、😍、🤩或😴。使用随包的授权图像，不依赖系统字体。保留原触发和时长，行走时与头顶留出间隙。"],
        ["音频","抽卡采用柔和提示；新英雄特写使用独立惊喜音。减少高频刺耳和十连叠音，服从全局音量开关。"],
        ["埋点","本期不新建分析后端。预留 home_enter、home_unlock、home_claim 的事件语义；真实发奖依据服务事务，日志不等于统计系统已上线。"],
    ],[1.4,5.25])
    doc.add_heading("10.1 关联系统与数据边界",level=2)
    body(doc,"家园沿用账号金币和券，居民读取卡牌；不改战斗金币、升级费用或抽卡概率。家园请求不信任自报价格、地块、时间或奖励；但现有整包客户端钱包提交机制仍存在，本次不是完整防作弊改造。原有阿里云测试服更新属于历史实施记录，本次仅修订客户端展示；Android 真机尚未验证。")
    doc.add_heading("10.2 配置失败与版本扩展",level=2)
    body(doc,"非法建筑类型、缺失物品或图标、负价格、无效圈数、零分母、概率超范围必须在配置验证阶段失败；运行时缺数据时禁用对应经济操作并明确原因。规则版本与已有 plot_id 价格保持稳定，后续调价、扩圈、存档迁移、拆迁退款必须另写迁移及补偿规则。")
    table(doc,["非本期扩展","当前边界"],[
        ["建筑升级与搬迁","无按钮、无价格、无隐式收益升级"],
        ["居民成长与需求条","纯观赏行为，不接饥饿、惩罚或付费恢复"],
        ["社交访问与排行榜","未实现，不共享他人家园可写状态"],
        ["长期生产与材料仓库","产物在本次日结中直接兑换，不建立可交易库存"],
    ],[2.1,4.55])

    page(doc,"11 验收条件与历史基线")
    body(doc,"以下为本次修订的验收要求，是否通过由 REQ-20260925-HOME-GACHA-READABILITY 回执与 temp/qa/home-readability-20260925 证据确认。历史测试单列保留，不证明本次规则已通过。本次不部署、不导出设备包；Windows 本地结果不代表 Android 真机结果。")
    table(doc,["ID","场景与通过条件","证据"],[
        ["H-Q01","真实鼠标/触摸点击左下入口打开家园；返回与重进正确","运行输入记录"],
        ["H-Q02","第一至第八圈价格范围、每圈 ×3 基准、目录稳定；价值单调并封顶","经济规则测试"],
        ["H-Q03","邻接价格不可读首击居中不扣款、次击建造；可读一击建造；非邻接无信息且不聚焦；解锁后邻接信息立即更新","事务与运行测试"],
        ["H-Q04","0/1/3/多于3天、UTC+8午夜、首次初始化、地块次日、同日重进","固定时间测试"],
        ["H-Q05","逐件金币范围、券概率公式和城堡保底；物品分页不改变次数","确定种子与边界测试"],
        ["H-Q06","重试、重复/并发请求、断线重查、保存失败回滚；不得重抽或二次入账","权威存储隔离测试"],
        ["H-Q07","游客/账号A/账号B状态隔离；切账号丢弃旧请求；重启恢复","持久化与账号回归"],
        ["H-Q08","无住宅全住城堡；容量1/5、超员、餐饮/运动白天、娱乐/休息夜晚","模拟行为测试"],
        ["H-Q09","所有路线段落在边路或门口过渡；可见出门、入楼隐藏、活动后气泡","几何断言与GPU截图"],
        ["H-Q10","720×1280、360×640；拖动各边缘不闪且不误买；最大值、禁用、次日产出无裁切重叠；道具图标数量品质一致","真实输入与运行PNG"],
        ["H-Q11","仅可领收益有红点；成功清除、失败保留、跨日恢复；可买地块不触发，切账号即时重算","运行断言"],
        ["H-Q12","GDScript解析与缩进、配置验证/导出、账户/卡牌/抽卡/战斗核心回归","Godot与测试日志"],
        ["H-Q13","Word真实标题、已更新TOC、全页无模板残留或缺字；图稿状态真实","DOCX包检查与PNG"],
    ],[0.65,4.25,1.75])
    doc.add_heading("11.1 历史验收基线",level=2)
    table(doc,["范围","检查数与失败数","证据与结论"],[
        ["家园历史运行基线","4626 / 0","home-runtime-final.log；Godot 4.6.2 Forward+，RTX 4060；存档替换、重开、重复领取与清理；双指干扰修复后通过"],
        ["家园经济","12381 / 0","economy.log；217 格、三天上限、次日产出、奖励与持久化"],
        ["同步与跨日","修复前 7 / 3；最终 12 / 0","home-sync-before.log、home-sync-final.log；慢回包保留新增资源；其他页签跨午夜也刷新家园红点"],
        ["相邻回归与兼容","奖励 50 / 0；战斗 10340 / 0","reward-regression.log、combat-regression.log；27 个旧 RPC 的签名顺序不变"],
        ["配置","19 表 467 行","验证与运行导出通过；CSV 和 JSON 同步"],
        ["玩家画面","720×1280 与 360×640","home-*.png；初始、扩建、昼夜、不足、解锁与三天奖励等独立状态"],
        ["价格可读后建造","修复前 10 / 3\n修复后 10 / 0","同一固定输入探针；价格不可读时首击居中不扣款，再次点击建造；证据见 v1.1 任务回执"],
    ],[1.3,1.7,3.65])

    page(doc,"12 图稿登记与评审状态")
    table(doc,["门禁","当前状态","责任人与关闭条件"],[
        ["用户需求与本地实施授权","已确认","用户当前请求；补充最多3天、解锁1券已合入"],
        ["规则与数据源文档","以最新修订为准","codex-primary；本文同步可读性规则，经济、配置和事务规则保持既有来源"],
        ["Word 全页渲染","本次独立复核","全页 PNG 与目录更新记录：temp/qa/home-readability-20260925/document；不代替运行验收"],
        ["正式 Figma 可编辑图稿","Pending","Figma whoami 发生 MCP HTTP 传输错误；工具恢复后创建、保存、对象回读和 PNG 核对"],
        ["运行逻辑与玩家可见验收","以本次回执为准","docs/receipts/REQ-20260925-HOME-GACHA-READABILITY.md；修复前后同探针、真实输入与GPU截图独立记录"],
        ["阿里云测试服家园更新","本次不涉及","原有部署记录保留；本次只修订客户端展示，不执行远程写入"],["生产服与设备发行","本次未执行","不属于本次可读性修订范围"],
    ],[1.6,1.45,3.6])
    doc.add_heading("12.1 Figma UI 与 UE 图稿登记",level=2)
    table(doc,["字段","值"],[
        ["工具选择","项目 AGENTS.md 指定 Figma/FigJam 优先；未把全局 Penpot 默认视为项目迁移许可"],
        ["figma_url / file_id / object_refs","Pending；未创建在线文件，不填虚构链接或对象 ID"],
        ["source_version / owner / review_state","v1.1 / codex-primary / Pending"],
        ["coverage","H00～H07；入口、主页、详情、产物、结算、失败重试、日夜与重进"],
        ["local_source_backup","tools/build_home_design_docx.py::build_figures；本文嵌入图按最新规则生成；原 docs/home_ui 图稿为历史草稿"],
        ["PNG 审阅导出","本文第3～4章嵌入图；本次图稿QA见 temp/qa/home-readability-20260925/document/figures"],
        ["机器登记","docs/home_ui/artifact_register.json 为原草稿登记，不代表本次在线图稿完成"],
    ],[2.2,4.45])
    body(doc,"文档和本地草稿支持当前已授权实现；外部图稿阻塞不自行撤销用户的开发指令。最后交付必须将文档、可编辑图稿、运行测试、线上部署四项分别说明，未取得的证据保持 Pending。")
    page(doc,"附录 抽卡首次获得动物演出")
    body(doc,"首次获得一种新动物时，先展示品质色卡背，点击或等待2秒后翻开，再展示大动物、名称、专属趣味台词和继续按钮。台词位于动物下方的对话框，只在特写阶段出现。无品质文字、属性、操作引导或倒计时。所有品质首得均触发，重复动物不触发；十连按原顺序展示，不改变概率、消耗或发卡事务。")
    table(doc,["页面与状态","显示与转换","数据和边界"],[
        ["G01 卡背与翻开","品质色卡背及光效；点击或2秒超时启动同一翻牌，完成后到G02","翻牌须可见收窄与展开，静止卡和美术等比；触发只一次，动画中屏蔽穿透"],
        ["G02 动物特写","正式动物美术、名称、下方趣味台词与继续按钮；继续到下一位G01","台词配置：cards.csv → cards.json 的 slogan。限32字、无换行；30px，最多两行，完整显示。空值隐藏，音效只播放一次"],
        ["G03 本次抽卡结果","十连2—3—3—2，横向40px、纵向28px间距，中间两排左右错开12px；动物等比缩为原适配尺寸78%","保持132×158卡片及24px名称，塔图不缩小；无品质、等级、升级进度。抽卡按钮整体移至底栏上方"],
        ["中断与连续输入","点击与超时竞争仅触发一次；切页清理旧计时与输入","奖励已由原事务写入；翻牌、连续点击、退出或重进不重复发卡或叠加音效"],
    ],[1.35,2.65,2.65])
    body(doc,"品质仅用卡牌、背景和光效颜色表达，不另写“绿色·普通”“蓝色·稀有”等文字。收藏与编组的玩法及升级信息保持原规则；精简只作用于抽卡结果和新动物演出。")
    doc.add_heading("附录 验收条件",level=2)
    for item in [
        "单抽首次获得：G01 使用正确品质色，点击与2秒超时都先播放翻牌再到G02，且只推进一次；卡牌下方无说明和倒计时。",
        "单抽重复卡：无首次获得演出，不改变原结果和卡牌数量；全动物已拥有的账号不能误标为首得。",
        "十连多个新动物：按结果顺序依次展示；同次抽到同种动物多张时只有第一次作为新种类揭晓。",
        "60种动物逐一检查配置台词；720×1280和360×640下台词完整且最多两行，与动物、名称和继续按钮不重叠。缺失、空值安全隐藏。",
        "十连按2—3—3—2错落排布，含阴影纵向净空至少22px；动物缩小22%，卡片和名称不缩小。无裁切、接触或遮挡按钮。",
        "关闭、连续点击、动画中切换与抽卡结果返回后无残留遮罩或计时，不阻塞下一次抽卡。",
    ]: add_bullet(doc,item)
    body(doc,"实现来源：scripts/app/ui/gacha_new_hero_reveal.gd 与 scripts/app/main.gd。本次需使用真实引擎帧计时和视口按下/释放验证翻牌，不替换抽卡或发卡逻辑。下表仅为旧演出的历史基线，不作为最新翻牌与精简布局的通过记录。")
    table(doc,["验收范围","结果与证据"],[
        ["原生输入集成","headless 45/0，GPU 61/0；gacha-qa.json；Godot 4.6.2 OpenGL 3.3 compatibility，RTX 4060"],
        ["全动物与全品质","60 种动物；展示逻辑 11736/0；gacha-reveal-logic.log"],
        ["16 张运行截图","4 品质 × 2 阶段 × 720/360 两尺寸；gacha-qa.json 逐图登记路径、大小与 SHA256"],
        ["范围限制","Windows 本地真实视口输入与渲染；未作移动设备导出验收；不新增网络或账号数据写入"],
    ],[1.65,5.0])
    doc.core_properties.title="家园模拟经营玩法设计"
    doc.core_properties.subject="F-ZC-HOME-001 v1.1 实现合同"
    doc.core_properties.author="Codex"
    doc.core_properties.keywords="家园,六边形,动物,每日收益,三天补领"
    return doc


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--figure-dir",type=Path,default=ART,help="Draft figure output directory; DOCX embeds these figures.")
    args=parser.parse_args()
    doc=build_document(args.figure_dir)
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))
    print("sha256="+hashlib.sha256(OUTPUT.read_bytes()).hexdigest())


if __name__ == "__main__":
    main()
