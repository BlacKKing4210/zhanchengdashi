from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"C:\Users\76398\.codex\skills\game-feature-design-docs\assets\general-feature-design-template.docx"
)
OUTPUT = ROOT / "docs" / "DEFENSE_TOWER_SYSTEM_DESIGN_v1.3.docx"

NAVY = "17243C"
BLUE = "2E75B6"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
GOLD = "F4B183"
LIGHT_GOLD = "FFF2CC"
RED = "C00000"
GRAY = "F2F2F2"
WHITE = "FFFFFF"

TOWERS = [
    ("defense_watch_tower", "兔子哨塔", 2, "绿色", 1, 5, 2.0, 1.5, "", "耳形屋脊＋单弩", "基础单体"),
    ("defense_longshot_tower", "猎鹰瞭望塔", 3, "蓝色", 1, 5, 3.5, 1.5, "攻击距离+1.5，优先攻击远程", "翼形檐口＋长望远镜", "远程优先"),
    ("defense_cannon_tower", "野猪重弩塔", 3, "蓝色", 2, 6, 2.0, 1.5, "攻击+1", "獠牙扶壁＋重型弩箭", "高单发"),
    ("defense_plunder_tower", "松鼠掠金塔", 3, "蓝色", 1, 7, 2.0, 1.5, "攻击时，掠夺1金币", "螺旋卷扬＋金币钩", "经济压制"),
    ("defense_rapid_tower", "麻雀速弩塔", 4, "紫色", 1, 6, 2.0, 0.5, "攻击速度+200%", "羽形檐片＋连发机括", "高频输出"),
    ("defense_repair_beacon", "龟甲壁垒塔", 4, "紫色", 1, 21, 2.0, 1.5, "生命值+200%", "龟甲穹顶＋大盾", "高耐久"),
    ("defense_twinshot_tower", "鹦鹉双弩塔", 4, "紫色", 1, 10, 3.0, 1.5, "攻击目标+1，攻击距离+1", "喙形遮檐＋双长弩", "双目标"),
    ("defense_bounty_tower", "老虎赏金塔", 5, "金色", 3, 12, 2.0, 1.5, "攻击+2，击杀时，获得10金币", "虎纹墙带＋重炮金币章", "击杀收益"),
    ("defense_territory_tower", "金雕领空塔", 5, "金色", 1, 12, 2.0, 1.5, "无视攻击距离，只要敌人处于我方领地上即可攻击", "金羽屋檐＋领地旗", "领地全域"),
    ("defense_storm_obelisk", "猛犸震地塔", 5, "金色", 1, 12, 2.0, 5.0, "每5秒对所有动物造成1点伤害", "象牙扶壁＋巨鼓震波", "全体脉冲"),
]


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text_color(cell, color: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def shade_row(row, color: str, text_color: str | None = None) -> None:
    for cell in row.cells:
        set_cell_fill(cell, color)
        if text_color:
            set_cell_text_color(cell, text_color)


def set_row_pagination(row, *, repeat_header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        table_header = OxmlElement("w:tblHeader")
        table_header.set(qn("w:val"), "true")
        tr_pr.append(table_header)


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in (
        ("Title", 28, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 11.5, NAVY),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("丛林法则｜F-ZC-DEFENSE-TOWER-005｜V1.3")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("6B7280")

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("内部制作评审｜2026-08-27")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("6B7280")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc: Document, text: str, *, bold: bool = False, color: str | None = None, size: float = 10.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(item)
        set_run_font(run, 10.5)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(f"{index}. {item}")
        set_run_font(run, 10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_row_pagination(header, repeat_header=True)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 9, True, WHITE)
        set_cell_margins(cell)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_pagination(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, WHITE if row_index % 2 == 0 else GRAY)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, 8.8)
            set_cell_margins(cell)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, 140, 160, 140, 160)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title + "\n")
    set_run_font(run, 10.5, True, RED if fill == LIGHT_GOLD else NAVY)
    run = paragraph.add_run(text)
    set_run_font(run, 10.2)
    doc.add_paragraph()


def add_title_page(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(92)
    run = paragraph.add_run("防御塔系统重制")
    set_run_font(run, 30, True, NAVY)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("最终属性、技能机制、描述与建筑化动物主题包装")
    set_run_font(run, 16, True, BLUE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    run = paragraph.add_run("F-ZC-DEFENSE-TOWER-005  ·  V1.3  ·  2026-08-27")
    set_run_font(run, 11, False, "6B7280")

    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata = [
        ("文档状态", "IMPLEMENTATION_CONTRACT；V1.3 数值与页面显示修订；方案 A 运行时美术保持不变"),
        ("正式来源", "制作人 10 塔表＋2026-08-27 全等级最终属性与页面隐藏距离/攻击间隔决策"),
        ("负责人", "制作人：user-producer｜实现：codex-primary"),
        ("目标运行环境", "Godot 4.6｜720×1280 竖屏"),
        ("数据口径", "表中攻击、生命、距离、间隔均为所有卡牌等级不再成长的效果后最终属性"),
    ]
    for index, (key, value) in enumerate(metadata):
        table.cell(index, 0).text = key
        table.cell(index, 1).text = value
        set_cell_fill(table.cell(index, 0), LIGHT_BLUE)
        for cell in table.rows[index].cells:
            set_cell_margins(cell, 130, 130, 130, 130)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, 10, index == 0)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    run = paragraph.add_run("重要：包装文字解释差异；运行时不得再次把文字中的加成叠到面板上。")
    set_run_font(run, 11.5, True, RED)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_version_control(doc: Document) -> None:
    add_heading(doc, "版本控制", 1)
    add_table(
        doc,
        ["版本", "编写人", "审核/批准", "日期", "更新内容"],
        [
            ["V1.0", "codex-primary", "user-producer", "2026-08-26", "10 塔最终属性、特殊技能、描述、UI 与整套美术包装合同"],
            ["V1.1", "codex-primary", "user-producer", "2026-08-27", "方案 B 曾进入生产准备，未绑定运行时即被后续制作人决策取代"],
            ["V1.2", "codex-primary", "user-producer", "2026-08-27", "改用方案 A；严格对齐已有动物画风，以大色块、粗轮廓和低细节生产并实装十塔"],
            ["V1.3", "codex-primary", "user-producer", "2026-08-27", "十塔四项属性在所有等级固定为制作人最终表；页面仅显示攻击、生命及原表技能描述"],
        ],
        [1.5, 2.6, 2.6, 2.4, 8.0],
    )
    add_table(
        doc,
        ["责任方", "负责人", "状态", "完成条件"],
        [
            ["策划", "user-producer", "已给定最终表", "数值及字面效果一致"],
            ["程序", "codex-primary", "V1.3 已实现并通过回归", "全等级属性锁定、两处页面隐藏和行为测试通过"],
            ["美术", "user-producer / codex-primary", "方案 A 运行时资产已验收，本次不改", "保留十张 480×480 RGBA 与现有绑定"],
            ["QA", "codex-primary", "V1.3 已验收", "配置、机制、720×1280 玩家可见证据与性能通过"],
        ],
        [2.2, 3.2, 3.0, 8.5],
    )


def add_design_content(doc: Document) -> None:
    add_heading(doc, "1. 术语缩写与修订标记", 1)
    add_table(
        doc,
        ["术语", "定义"],
        [
            ["最终属性", "制作人表中的攻击、生命、距离（格）、攻击间隔（秒）；已包含包装效果，且防御塔卡在所有等级均保持该值。"],
            ["技能机制", "由稳定 skill_id 驱动的目标选择、金币转移、多目标、击杀奖励、领地目标或全体脉冲。"],
            ["包装效果", "玩家用于理解塔定位的直接文案；除明确技能机制外，不再修改最终基础属性。"],
            ["我方领地", "目标动物当前位置所在格，其可视归属阵营与塔所属阵营为同盟。"],
        ],
        [4.0, 12.5],
    )

    add_heading(doc, "2. 设计目的", 1)
    add_heading(doc, "2.1 主要目标", 2)
    add_bullets(doc, [
        "把制作人给定的 10 座塔一一落到卡牌数据、战斗属性、特殊行为和技能描述。",
        "玩家在卡牌详情和战斗建筑详情中只看到攻击、生命及制作人原表技能描述，不显示距离和攻击间隔。",
        "保留旧存档依赖的 4 个稳定塔 ID，新增 6 个塔 ID，避免账号与 AI 卡组失效。",
        "按已批准方案 A 建立 10 座建筑优先、动物同风格的极简包装：使用大色块、粗深色轮廓和少量平涂阴影，并用真实运行切片证明卡牌与战场可读性。",
    ])
    add_heading(doc, "2.2 次要目标", 2)
    add_bullets(doc, [
        "清除旧代码统一加半格射程、固定 1 秒间隔造成的二次计算。",
        "统一 cards、defenses、skills、localization 的策划口径。",
        "延续塔的粘性锁敌，避免频繁重选造成观感和性能回退。",
    ])
    add_heading(doc, "2.3 非目标", 2)
    add_bullets(doc, [
        "本版本不调整塔的抽卡品质总概率。",
        "本版本不改变动物、金矿及其他非防御塔卡的等级成长；仅防御塔四项战斗属性在所有等级固定为制作人表。",
        "本版本不改变塔的碰撞、占地、生命条数据或未解锁塔的通用剪影规则。",
        "不新增塔升级树、弹药、主动施法或新的战斗页面。",
    ])

    add_heading(doc, "3. 功能概述", 1)
    add_numbered(doc, [
        "玩家获得并把防御塔卡放入编组；绿色兔子哨塔仍是初始强制塔。",
        "战斗解锁塔地块后，系统按该塔卡的制作人最终属性创建建筑；卡牌等级不放大塔的攻击、生命、距离或攻击间隔。",
        "塔按攻击间隔结算；普通塔单体攻击，特殊塔通过 skill_id 改变选敌或结算。",
        "卡牌详情与战斗中的塔详情仅显示攻击、生命及完整技能描述；基础塔原表技能为空时不显示技能行。",
        "方案 A 的动物同风格低细节正式资源通过代表塔运行切片后，每座已建成塔按 site_card -> art_path 显示唯一建筑化动物主题 PNG；缺图或旧存档回退通用塔。",
    ])
    add_callout(
        doc,
        "最终属性硬规则",
        "表中攻击、生命、距离与攻击间隔已经考虑“+1”“+200%”“+1格”等效果。十塔在任意卡牌等级都只读取这些最终值，禁止等级倍率或包装文案再次加成。",
    )

    add_heading(doc, "4. 系统框架", 1)
    add_table(
        doc,
        ["层", "正式来源", "职责", "禁止项"],
        [
            ["卡牌数据", "config/tables/cards.csv", "运行时基础面板、skill_id、技能描述、卡图路径", "不从中文文本推断技能"],
            ["塔策划镜像", "config/tables/defenses.csv", "面向关卡/工具的同口径塔数据", "不得与 cards 冲突"],
            ["技能目录", "config/tables/skills.csv", "技能分类、目标、强度、说明键", "不替代实际 GDScript 结算"],
            ["纯规则", "DefenseTowerRules", "按 skill_id 输出行为开关和数值", "不持有场景状态"],
            ["战斗编排", "scripts/app/main.gd", "目标合法性、伤害、金币、反馈、权威端结算", "不二次加面板属性"],
            ["玩家显示", "cards.skill_text", "卡牌详情与战斗塔详情的唯一可见描述", "不作为逻辑条件"],
        ],
        [2.6, 4.2, 6.0, 4.0],
    )

    add_heading(doc, "5. UE 总流程与页面覆盖", 1)
    add_callout(
        doc,
        "可编辑 UE 源状态",
        "本功能不新增页面，只扩展既有收藏/编组卡牌详情与战斗建筑详情。当前环境无可用 Figma/FigJam 连接器，因此不伪造可编辑链接；页面行为以下表为正式文字合同，美术/UE 可编辑源标记为待补。",
        fill=LIGHT_BLUE,
    )
    add_table(
        doc,
        ["页面/状态", "入口", "玩家动作/事件", "反馈", "返回/异常"],
        [
            ["P-01 收藏/编组卡牌详情", "点击任意防御塔卡", "查看属性与技能描述", "只显示攻击、生命及原表技能；不显示距离/攻击间隔", "技能为空不显示技能行；其余文案无截断"],
            ["P-02 战斗塔详情", "点击已建成防御塔", "查看塔属性与技能描述", "3 秒详情条；只显示攻击、生命及原表技能", "塔被摧毁或换卡立即关闭"],
            ["S-01 塔攻击循环", "spawn_timer≤0", "系统选敌并结算", "弹道/脉冲、伤害、金币反馈", "无合法目标则本次空放并进入下一间隔"],
            ["S-02 代表运行切片", "方案 A 鹦鹉双弩塔动物同风格低细节正式图", "打开卡牌详情并在战场建成", "720×1280 同时验证双弩、建筑感、大色块、详情与血条", "BLOCKER/MATERIAL 必须返修"],
            ["S-03 十塔全量绑定", "代表切片通过", "生成整套、标准化、写入 art_path", "已建成塔显示唯一图；未解锁仍是通用剪影", "缺图/空 ID 回退通用塔"],
        ],
        [3.2, 3.2, 4.0, 4.2, 3.0],
    )

    add_heading(doc, "6. 正式数值与包装表", 1)
    rows = []
    for tower_id, name, tier, color, attack, hp, distance, interval, skill_text, theme, role in TOWERS:
        rows.append([
            str(tier),
            color,
            name,
            str(attack),
            str(hp),
            ("%g格" % distance),
            ("%gs" % interval),
            skill_text,
            role,
        ])
    add_table(
        doc,
        ["品质", "品质色", "防御塔", "攻", "血", "距离", "间隔", "技能描述", "定位"],
        rows,
        [1.0, 1.2, 2.5, 0.8, 0.8, 1.2, 1.2, 6.2, 1.8],
    )
    add_text(doc, "注：所有行均为效果后最终面板，防御塔卡从 Lv.1 到最高等级都必须严格保持这些数值；品质 2/3/4/5 分别映射 common/rare/epic/legendary。", bold=True, color=RED)

    add_heading(doc, "7. 配置表调整", 1)
    add_table(
        doc,
        ["策划含义", "字段", "类型/单位", "规则"],
        [
            ["塔稳定 ID", "cards.csv::id / defenses.csv::id", "id", "保留 4 个旧 ID，新增 6 个"],
            ["最终攻击", "cards.csv::attack", "int", "所有等级直接读取，不叠加包装或等级倍率"],
            ["最终生命", "cards.csv::max_hp", "int", "建造时写入 hp/max_hp"],
            ["最终距离", "cards.csv::attack_range", "float/格", "战斗判定转换为 格×HEX_SIZE"],
            ["最终间隔", "cards.csv::summon_interval_sec", "float/秒", "允许 0.5 秒，不再强制 1 秒"],
            ["机制 ID", "cards.csv::skill_id", "id", "唯一运行时技能分派键"],
            ["显示文案", "cards.csv::skill_text", "string", "玩家可见唯一来源；逐字显示；空值不绘制技能行"],
            ["技能目录", "skills.csv::*", "多字段", "分类与工具引用；行为仍由稳定 ID 实现"],
            ["中文文本", "localization_zh_runtime.csv::*", "string", "塔名及技能描述键"],
        ],
        [3.0, 5.0, 2.2, 6.5],
    )

    add_heading(doc, "8. 系统逻辑", 1)
    add_heading(doc, "8.1 总体结算", 2)
    add_numbered(doc, [
        "权威战斗端读取塔卡最终面板，不应用卡牌等级倍率；距离字段只在内部从格转换为世界单位。",
        "若为猛犸震地塔，直接执行全体脉冲并结束本次攻击，不再选单体目标。",
        "其他塔先验证粘性锁定；无合法锁定时，按技能策略选择新目标。",
        "对主目标执行一次完整面板攻击；随后结算掠夺、额外目标或击杀奖励。",
        "重置该塔 spawn_timer 为制作人表中的最终攻击间隔；非权威客户端只应用快照。",
    ])

    add_heading(doc, "8.2 特殊机制", 2)
    add_table(
        doc,
        ["skill_id", "触发", "精确定义", "关键边界"],
        [
            ["defense_far_sight", "获取新目标", "射程内有远程动物时，只在远程动物中选最近；否则按普通最近。", "合法粘性目标不因新远程目标出现而换靶。"],
            ["defense_plunder", "成功直接攻击后", "从主目标所属阵营转移 min(1, 目标当前金币) 给塔阵营。", "盟友/中立不触发；金币总量守恒。"],
            ["defense_twin_shot", "主攻击后", "在范围内选择一个不同的敌方动物，造成完整面板攻击。", "不重复主目标；最多额外 1 个；不额外攻击建筑。"],
            ["defense_bounty", "实际击杀后", "本次攻击每实际击杀 1 只动物，塔阵营获得 10 金币。", "非致死、已死亡目标、建筑摧毁均不触发。"],
            ["defense_territory", "验证/获取目标", "任意距离的敌方动物，只要当前位置地块属于塔阵营或其盟友，即为合法。", "中立、虚空、敌方领地和建筑均非法。"],
            ["defense_global_pulse", "每 5 秒", "对触发开始时的全部存活动物各造成 1 点伤害。", "含己方/盟友/敌方；不伤建筑；替代普通单体攻击。"],
        ],
        [3.3, 2.5, 7.0, 4.2],
    )

    add_heading(doc, "8.3 状态与边界矩阵", 2)
    add_table(
        doc,
        ["场景", "系统处理", "玩家反馈", "验收"],
        [
            ["塔无目标", "不造成伤害，继续下一间隔", "无弹道", "不报错、不锁无效 ID"],
            ["被锁目标死亡/离开合法区", "清锁并按策略重选", "下一次攻击转向", "粘性锁敌保持"],
            ["掠夺目标 0 金币", "转移 0", "无金币增加反馈", "金币不凭空生成"],
            ["双射仅有一个合法动物", "只打主目标", "一次弹道", "不重复命中"],
            ["领地归属变化", "下一次验证立即失效或生效", "正常换靶", "读取实时位置地块"],
            ["脉冲中有动物死亡", "对启动时快照逐个结算；新召唤动物不进入本轮", "全体受击反馈", "每只最多 1 点"],
            ["多人非权威客户端", "不运行塔计时与技能", "只显示服务器快照", "无双重结算"],
        ],
        [3.2, 6.0, 4.2, 4.0],
    )

    add_heading(doc, "9. UI 界面及显示规则", 1)
    add_table(
        doc,
        ["元素", "显示规则", "数据源", "异常/空状态"],
        [
            ["卡牌塔图", "显示方案 A 唯一 480×480 RGBA 动物同风格低细节塔图，保持等比居中", "cards.csv::art_path", "缺图回退通用 tower.png，不得显示兔子兜底"],
            ["品质", "沿用绿色/蓝色/紫色/金色卡框", "cards.csv::rarity", "无"],
            ["两项可见属性", "仅显示攻击、生命；页面不得显示距离或攻击间隔", "防御塔最终 stats", "内部仍保留距离/间隔用于战斗"],
            ["技能描述", "严格显示制作人原表；长文最多两行，自适应字号，不显示省略号", "cards.csv::skill_text", "空值不显示技能行，不自行补写说明"],
            ["战斗塔详情", "点击已建成塔显示 3 秒；结构复用动物营地详情条", "tile.site_card", "塔摧毁/卡不匹配立即关闭"],
        ],
        [3.0, 6.0, 4.3, 4.0],
    )

    add_heading(doc, "10. 相关需求", 1)
    add_heading(doc, "10.1 美术资源需求", 2)
    art_rows = []
    for index, (_, name, tier, color, _attack, _hp, _distance, _interval, _text, theme, _role) in enumerate(TOWERS, start=1):
        art_rows.append([f"T-{index:02d}", name, theme, f"品质{tier}/{color}", "方案 A；动物同风格大色块低细节；480×480 RGBA；透明；建筑基线 y=438±4；运行时绑定"])
    add_table(doc, ["ID", "名称", "唯一记忆点", "品质", "规格/状态"], art_rows, [1.2, 3.0, 4.8, 2.2, 6.5])
    add_bullets(doc, [
        "风格：与现有动物 PNG 一致的紧凑 Q 版平面插画；粗深色墨线、大面积纯色或两档平涂、极少纹理、透明背景；第一眼必须是可工作的防御建筑。",
        "复杂度预算：每塔只用 3～5 个主形状、4～6 个稳定色彩角色、2～4 块大面积墙体分区；禁止逐块砖缝、微小铆钉、密集羽毛线、金币纹理、碎旗和装饰噪点。",
        "构图：塔体占主体轮廓至少 85%；动物元素仅建筑化为屋脊、檐口、扶壁、甲片、纹样、机括或导管；每塔只有一个主武器/技能道具，双弩塔可用一对同构武器。",
        "禁用：完整动物站、坐、骑乘、驾驶或停驻在塔上；独立巨大动物头；眼鼻嘴构成的脸形；文字、数字、UI 框、场景背景、渐变光效与碎小装饰。",
        "流程：方案 A 方向批准 → 鹦鹉双弩塔动物同风格低细节代表图 → 720×1280 代表运行切片 → 其余九塔 → 5×2 整板冻结像素与哈希 → 十塔全量绑定。",
        "透明规范：480×480 RGBA8；视觉中心 x=240±6；建筑基线 y=438±4；四角至少 24×24 全透明；不保留白底、格线、文字或烘焙地面阴影。",
        "脸形禁用：不得以眼点、鼻孔、喙、象鼻或嘴形成动物面部关系；鹦鹉双弩塔的金色三角件只可读作无眼雨棚/箭槽。",
    ])

    add_heading(doc, "10.2 音乐音效需求", 2)
    add_text(doc, "本期不新增音频资源。普通攻击沿用 tower_attack；金币沿用 gold_gain；全体脉冲沿用受击与脉冲反馈。后续如批准独立音效，另立资源批次。")
    add_heading(doc, "10.3 功能打点需求", 2)
    add_table(
        doc,
        ["事件", "时机", "参数", "目的"],
        [
            ["tower_attack_resolved", "每次塔攻击完成", "tower_id, team, target_count, damage", "核验命中与性能"],
            ["tower_gold_transferred", "掠夺完成", "tower_id, from_team, to_team, amount", "核验经济守恒"],
            ["tower_kill_bounty", "悬赏发放", "tower_id, team, killed_unit_id, amount", "核验击杀奖励"],
            ["tower_global_pulse", "全体脉冲", "tower_id, unit_count, damage", "性能与影响面"],
        ],
        [4.0, 3.5, 6.0, 3.5],
    )

    add_heading(doc, "11. 关联拓展", 1)
    add_bullets(doc, [
        "关联卡牌：初始强制绿色塔、抽卡、收藏、编组、升级、段位 AI 卡组。",
        "关联战斗：建筑创建、塔计时、目标锁定、伤害、金币、领地归属、多人权威快照。",
        "关联美术：现有 60 动物风格合同；正式塔图不得复制外部商业游戏资产或识别。",
        "后续可扩展：塔技能 VFX、独立音效、塔升级外观或非属性升级收益；均不在 V1.3 范围。",
    ])

    add_heading(doc, "12. 验收与 QA", 1)
    add_heading(doc, "12.1 交付验收", 2)
    qa_rows = [
        ["数据", "10 个稳定 ID；品质与任意等级攻/血/距离/间隔逐行等于正式表", "自动测试＋运行时配置导出"],
        ["不二次计算", "无统一半格射程；无固定 1 秒；0.5/1.5/5 秒均可生效", "单元测试"],
        ["远程优先", "新目标优先远程；合法旧目标保持", "行为测试"],
        ["掠夺", "最多转移 1；目标为 0 不产币；总量守恒", "行为测试"],
        ["双射", "恰好最多 2 个不同敌方动物；完整攻击伤害", "行为测试"],
        ["悬赏", "只对实际击杀动物发 10；建筑和非致死不发", "行为测试"],
        ["领地", "领地内任意距离可攻击；领地外/中立/虚空/建筑不可", "行为测试"],
        ["全体脉冲", "每 5 秒替代普攻；触发时所有存活动物各 1 点", "行为＋性能测试"],
        ["页面显示", "卡牌详情与战斗塔详情均只显示攻/血；不显示距离/间隔；基础塔无技能行，其余9塔原文完整", "UI合同测试＋720×1280 GPU 截图"],
        ["美术", "恰好 10 张 480×480 RGBA；方案 A 塔体主体≥85%；动物同风格大色块低细节；无完整动物/巨大动物头；Alpha、基线、小图与整板检查通过", "生产 manifest＋浅/深底＋48/56/64/96px 预览"],
        ["运行时绑定", "10 个 art_path 唯一可加载；已建成塔用 site_card 图；未解锁/缺图/空 ID 安全回退通用塔", "Godot 专项测试＋720×1280 GPU 截图"],
        ["回归", "配置、缩进、Godot 解析、塔卡组、AI、多人权威通过；72 单位 P95＜4ms", "自动回归＋运行证据"],
    ]
    add_table(doc, ["验收域", "通过条件", "证据"], qa_rows, [2.4, 10.2, 4.5])

    add_heading(doc, "12.2 待确认与评审记录", 2)
    add_table(
        doc,
        ["问题/决策", "状态", "责任人", "结论及影响"],
        [
            ["10 塔建筑化动物包装方向", "已批准方案 A；方案 B 已被取代", "user-producer", "对齐现有动物画风；大色块、粗轮廓、低细节；v2 方向板仍为 NOT_RUNTIME"],
            ["可编辑 Figma/FigJam UE 源", "工具不可用/待补", "user-producer", "不新增页面；文字合同不受影响"],
            ["所有动物是否包含己方", "按字面已实现", "user-producer", "V1.0 包含全部阵营动物"],
            ["新增塔对同品质抽卡池的稀释", "沿用现有机制", "user-producer", "本期不改品质总概率"],
            ["防御塔卡升级入口", "保留现状", "user-producer", "等级仍可提升，但 V1.3 不再提供四项战斗属性成长；非属性收益另案设计"],
        ],
        [6.3, 3.0, 3.2, 5.0],
    )


def build() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    set_document_defaults(doc)
    doc.core_properties.title = "防御塔系统重制 V1.3"
    doc.core_properties.subject = "10座防御塔全等级最终属性、技能描述与页面显示合同；方案A运行时美术保持不变"
    doc.core_properties.author = "codex-primary"
    doc.core_properties.keywords = "Godot, 防御塔, 技能, 建筑化动物包装, F-ZC-DEFENSE-TOWER-005"
    add_title_page(doc)
    add_version_control(doc)
    add_design_content(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
