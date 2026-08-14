from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\76398\.codex\skills\game-feature-design-docs\assets\general-feature-design-template.docx")
OUTPUT = ROOT / "docs" / "RENDERING_CLARITY_AND_BATTLE_PERFORMANCE_DESIGN_v1.0.docx"

INK = "172033"
BLUE = "2E74B5"
TEAL = "197A78"
GRAY = "5B6574"
LIGHT_BLUE = "EAF3FA"
LIGHT_GOLD = "FFF4D6"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Microsoft YaHei"
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    set_run_font(paragraph.add_run(text), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text))


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text))


def add_body(doc: Document, text: str, *, color: str = INK, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(text), bold=bold, color=color)


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 12.5, TEAL, 10, 5),
        ("Heading 3", 11, GRAY, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.text = "丛林法则｜渲染清晰度与战斗性能技术合同"
    set_run_font(header.runs[0], 8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "v1.0｜2026-08-14｜IMPLEMENTATION_CONTRACT"
    set_run_font(footer.runs[0], 8.5, color=GRAY)


def add_title_page(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(6)
    set_run_font(title.add_run("渲染清晰度与战斗性能优化"), 24, True, INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("移动端竖屏 1080×1920 目标｜UI 原生像素栅格化｜72 单位性能基准"), 12.5, False, BLUE)

    add_table(
        doc,
        ["字段", "内容"],
        [
            ["文档类型", "IMPLEMENTATION_CONTRACT｜MATERIAL"],
            ["请求编号", "REQ-20260814-RENDER-CLARITY-BATTLE-PERF"],
            ["功能编号", "F-ZC-001"],
            ["版本 / 日期", "v1.0 / 2026-08-14"],
            ["负责人", "制作人：用户｜工程负责人：codex-primary"],
            ["目标平台", "Android 竖屏；Windows 开发验证"],
            ["正式目标", "1080×1920；540×960 作为低分辨率兼容验证"],
            ["体验记忆点", "同一布局下，标题和资源数字清晰，满单位战斗仍保持连贯响应"],
            ["当前门禁", "PROTOTYPE_READABLE → RUNTIME_SLICE_APPROVED（完成运行时证据后）"],
        ],
        [1.45, 5.1],
    )

    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_cell_shading(callout.cell(0, 0), LIGHT_GOLD)
    set_cell_text(
        callout.cell(0, 0),
        "制作人决策：不改变页面信息、布局、按钮位置、点击区域、状态含义和战斗镜头构图；只提高有效渲染清晰度并降低战斗每帧开销。",
        bold=True,
        size=10.5,
    )
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    for line in (
        "1. 术语缩写与修订标记",
        "2. 设计目的",
        "3. 功能概述",
        "4. 系统框架",
        "5. UE 总流程与页面覆盖",
        "6. 参考与基线",
        "7. 配置与数据源调整",
        "8. 系统逻辑",
        "9. UI 界面及状态",
        "10. 相关需求",
        "11. 关联拓展",
        "12. 验收与 QA",
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(line), 10, color=GRAY)
    doc.add_page_break()


def build_document() -> Document:
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    style_document(doc)
    add_title_page(doc)
    add_toc(doc)

    doc.add_heading("1. 术语缩写与修订标记", level=1)
    add_table(
        doc,
        ["术语", "定义"],
        [
            ["逻辑分辨率", "现有 720×1280 设计坐标；用于布局和点击命中，不等同于实际屏幕像素。"],
            ["原生像素栅格化", "先把逻辑文字矩形映射到实际视口，再按实际像素字号生成字形，避免缩放小字形贴图。"],
            ["p95", "采样值从小到大排列后第 95 百分位；用于反映偶发慢帧。"],
            ["单位索引缓存", "每次模拟更新构建一次 unit_id → 数组下标映射，并在使用前验证。"],
        ],
        [1.55, 5.0],
    )
    add_table(
        doc,
        ["版本", "日期", "变更", "状态"],
        [["v1.0", "2026-08-14", "建立清晰度、分辨率与战斗性能实施合同", "制作人已提出实施"]],
        [0.8, 1.1, 3.7, 0.95],
    )

    doc.add_heading("2. 设计目的", level=1)
    doc.add_heading("2.1 主要目标", level=2)
    add_bullet(doc, "标题、资源数字、按钮文字、底部导航文字在 540×960 与 1080×1920 实际视口上均按原生像素字号绘制，显著减少模糊边缘。")
    add_bullet(doc, "UI 细线与边框使用 2D 像素对齐，保持现有视觉层级和触摸命中范围。")
    add_bullet(doc, "72 单位锁定目标场景下，减少平方级重复查询与单位循环内部的重复建筑索引重建，降低模拟慢帧。")
    doc.add_heading("2.2 次要目标", level=2)
    add_bullet(doc, "建立可重复的双分辨率截图和固定单位数基准，后续每次改 UI 或战斗规则都能复测。")
    add_bullet(doc, "在不增加 2K/4K 固定离屏渲染开销的前提下，提高有效清晰度。")
    doc.add_heading("2.3 非目标", level=2)
    add_bullet(doc, "不重做 UI 美术、不移动控件、不修改文字内容、不更换玩法流程、不调整战斗数值和单位上限。")
    add_bullet(doc, "本版本不以 Windows 基准替代 Android 真机验收；没有代表性手机证据时，不宣称移动端绝对无卡顿。")

    doc.add_heading("3. 功能概述", level=1)
    add_number(doc, "游戏仍以 720×1280 逻辑坐标计算页面布局和点击区域。")
    add_number(doc, "每次绘制时把文字矩形映射到当前实际视口，按 scale 后的整数像素字号生成字形，并在绘制后恢复原变换。")
    add_number(doc, "项目维持 1080×1920 竖屏目标；540×960 屏幕直接生成匹配该屏幕的字形，不先生成 720 基准字形再缩小。")
    add_number(doc, "战斗每个模拟更新先生成单位 ID 索引和建筑目标索引，后续单位复用；锁定目标查询优先 O(1) 命中。")
    add_number(doc, "自动化同时输出双分辨率截图、清晰度结构断言以及 72 单位 median/p95 基准。")

    doc.add_heading("4. 系统框架", level=1)
    add_table(
        doc,
        ["层", "输入", "处理", "输出 / 约束"],
        [
            ["逻辑布局", "720×1280 页面坐标", "沿用 MainPageLayout 与现有 Rect2", "页面结构、点击区域完全不变"],
            ["视口映射", "实际 viewport_size", "scale=min(W/720,H/1280)，offset 居中", "支持 9:16 与更长屏幕留白背景"],
            ["文字渲染", "逻辑 rect、字号、当前 draw transform", "映射到屏幕像素，字号取 max(1, round(size×scale))", "清晰字形；绘制后恢复世界/页面变换"],
            ["UI 几何", "现有 draw_rect / draw_line", "启用 2D transform/vertex pixel snap", "细边框减少半像素抖动"],
            ["战斗索引", "units、tiles", "每模拟更新构建一次索引", "锁定目标查询 O(1)，建筑键只重建一次"],
            ["QA", "固定场景与同一 Godot 4.6.2", "双截图 + 72 单位计时", "可复现输出，不用肉眼主观替代数据"],
        ],
        [1.0, 1.35, 2.6, 1.75],
    )

    doc.add_heading("5. UE 总流程与页面覆盖", level=1)
    add_body(doc, "本任务不新增页面、弹窗、入口、分支或交互状态，现有 UE 总流程保持不变，因此不创建新的 Penpot 流程节点。运行路径仍为：大厅 → 单人/多人战斗 → 战斗 HUD → 暂停/结算 → 返回大厅。")
    doc.add_heading("5.1 页面 UE 图与总流程节点覆盖表", level=2)
    add_table(
        doc,
        ["页面 / 状态", "变化", "必须保持"],
        [
            ["大厅", "文字与边框清晰度", "资源条、模式按钮、底部导航位置与点击区"],
            ["战斗 HUD", "文字与图形清晰度；模拟性能", "镜头、棋盘、资源条、暂停键、选中面板"],
            ["房间 / 账号输入", "继承原生像素字号与像素对齐", "LineEdit 触摸聚焦、软键盘和提交行为"],
            ["暂停 / 结算", "继承原生像素字号", "按钮含义、顺序、返回逻辑"],
        ],
        [1.35, 2.15, 3.2],
    )

    doc.add_heading("6. 参考与基线", level=1)
    add_table(
        doc,
        ["来源", "证据", "用途"],
        [
            ["制作人截图", "540×960 PNG；SHA-256 E55A66D5…1E81F0F4", "识别标题、资源值、按钮和导航文字发糊"],
            ["project.godot", "viewport 1080×1920；canvas_items；expand；portrait", "确认声明分辨率已经高于实际截图"],
            ["scripts/app/main.gd", "DESIGN_SIZE 720×1280；统一 CanvasItem scale", "确认小字号字形随画布缩放的根因"],
            ["固定 72 单位基准", "同一 Godot 4.6.2、同一脚本、同一迭代数", "比较优化前后 median/p95"],
        ],
        [1.35, 3.45, 1.9],
    )

    doc.add_heading("7. 配置与数据源调整", level=1)
    add_table(
        doc,
        ["来源", "字段 / 常量", "值 / 规则", "说明"],
        [
            ["project.godot", "display/window/size/viewport_width", "1080", "正式竖屏目标宽度，保持"],
            ["project.godot", "display/window/size/viewport_height", "1920", "正式竖屏目标高度，保持"],
            ["project.godot", "rendering/2d/snap/snap_2d_transforms_to_pixel", "true", "对齐 2D 变换"],
            ["project.godot", "rendering/2d/snap/snap_2d_vertices_to_pixel", "true", "对齐 2D 顶点"],
            ["main.gd", "DESIGN_SIZE", "720×1280", "保持布局和输入兼容"],
            ["main.gd", "native_font_size", "max(1, round(logical_size×effective_scale))", "按实际像素生成字形"],
        ],
        [1.35, 2.45, 1.5, 1.4],
    )
    add_body(doc, "本任务不调整 config/tables/*.csv，不触发 CSV 导出。", color=GRAY)

    doc.add_heading("8. 系统逻辑", level=1)
    doc.add_heading("8.1 原生像素文字", level=2)
    add_bullet(doc, "所有公共文字辅助函数先用逻辑字号执行截断与排版，保证原有文本宽度和布局不变。")
    add_bullet(doc, "根据当前设计/世界绘制变换，把逻辑矩形映射成屏幕矩形；字号按当前均匀缩放取整。")
    add_bullet(doc, "临时切换为单位屏幕变换绘制字形，完成后恢复页面或棋盘世界变换，保证地块费用等世界文字位置正确。")
    add_bullet(doc, "LineEdit 等引擎原生 Control 延续现有按 canvas_scale 设置字号的路径，不重复缩放。")
    doc.add_heading("8.2 战斗查询缓存", level=2)
    add_bullet(doc, "每次 _update_units 开始时重建 unit_id → index；命中时校验数组范围和 ID，一旦不一致立即线性回退并修复缓存。")
    add_bullet(doc, "死亡单位在本帧末尾压缩后再次重建缓存；本帧追加单位不会破坏已有下标，首次查询可安全回退。")
    add_bullet(doc, "combat_building_keys 在每次单位更新开始时生成一次；_ensure_unit_navigation_target 不再为每个无目标单位重复生成。")
    doc.add_heading("8.3 状态与边界矩阵", level=2)
    add_table(
        doc,
        ["状态", "预期行为", "失败保护"],
        [
            ["540×960", "scale=0.75；字号按实际像素取整", "最小字号 ≥1；布局仍居中"],
            ["1080×1920", "scale=1.5；字形以 1.5 倍实际字号生成", "不修改逻辑 Rect2"],
            ["更长竖屏", "统一 scale，剩余区域用全屏背景填充", "canvas_offset 居中"],
            ["单位新增", "缓存未含新 ID 时线性回退并写回", "禁止返回错误下标"],
            ["单位死亡压缩", "压缩后立即重建缓存", "选中单位失效检查继续执行"],
            ["无建筑目标", "建筑键为空，单位保持安全状态", "不重复全图扫描"],
        ],
        [1.25, 3.1, 2.25],
    )

    doc.add_heading("9. UI 界面及状态", level=1)
    doc.add_heading("9.1 页面清单", level=2)
    add_body(doc, "大厅、编组、抽卡、房间、战斗、账号中心、暂停、结算全部继承同一文字渲染辅助函数。没有新页面。")
    doc.add_heading("9.2 页面元素规范", level=2)
    add_bullet(doc, "标题和主 CTA：保持字号、颜色、矩形和对齐方式；只改变字形生成分辨率。")
    add_bullet(doc, "资源条和底部导航：保持图标、数值槽、边框宽度与触摸区域；启用像素对齐。")
    add_bullet(doc, "战斗棋盘与单位：保持镜头 1.30、动物 35% 集成缩放和血条位置；本任务不改变单位视觉大小。")
    add_bullet(doc, "文字截断：继续使用现有省略号规则；屏幕字号取整不能导致逻辑宽度重新排版。")

    doc.add_heading("10. 相关需求", level=1)
    doc.add_heading("美术资源需求", level=2)
    add_body(doc, "无新美术资源。现有 PNG 不做 AI 重绘或重采样覆盖；清晰度来自运行时渲染链路。")
    doc.add_heading("音乐音效需求", level=2)
    add_body(doc, "无。")
    doc.add_heading("功能打点需求", level=2)
    add_table(
        doc,
        ["事件 / 指标", "字段", "用途"],
        [
            ["开发基准输出", "viewport、unit_count、warmup、samples、median_us、p95_us", "回归比较；仅测试日志，不进入玩家日志"],
            ["截图证据", "viewport、scale、native_font_size、path", "验证双分辨率渲染"],
        ],
        [1.55, 3.2, 1.85],
    )

    doc.add_heading("11. 关联拓展", level=1)
    add_bullet(doc, "后续可在代表性 Android 低端、中端设备上记录 GPU/CPU Profiler 和帧时间，再决定是否引入更强的空间分区；本版本不提前复杂化。")
    add_bullet(doc, "若未来更换为可分发 CJK 字体资源，应单独完成字体授权、字形覆盖、包体积和 Android 真机检查。")
    add_bullet(doc, "如果页面布局或视觉语言发生变化，必须另走 editable Penpot 页面与视觉质量合同；本任务不提供该授权。")

    doc.add_heading("12. 验收与 QA", level=1)
    doc.add_heading("12.1 交付验收", level=2)
    add_table(
        doc,
        ["ID", "验收项", "通过标准"],
        [
            ["RC-01", "项目分辨率", "1080×1920、portrait、canvas_items/expand 保持；2D pixel snap 开启"],
            ["RC-02", "540×960 清晰度", "截图存在；标题、资源数字、按钮、导航使用实际像素字号"],
            ["RC-03", "1080×1920 清晰度", "截图存在；逻辑布局与 540×960 一致；字形按 1.5 倍字号生成"],
            ["RC-04", "变换恢复", "世界地块费用文字与页面文字位置正确，无漂移或缩放串扰"],
            ["BP-01", "72 单位性能", "同机同脚本 optimized p95 < 4.0 ms，且不劣于 baseline"],
            ["BP-02", "战斗规则回归", "经典战斗、单位检查、动物技能、多人目标相关测试无新增失败"],
            ["QA-01", "代码质量", "GDScript 缩进检查、项目解析、staged-tree 验证通过"],
            ["QA-02", "移动端边界", "无真机证据时状态明确为 Android device acceptance pending"],
        ],
        [0.7, 2.05, 3.85],
    )
    doc.add_heading("12.2 待确认与评审记录", level=2)
    add_table(
        doc,
        ["事项", "当前状态", "责任 / 下一步"],
        [
            ["制作人视觉确认", "待运行时双分辨率截图", "制作人检查文字与 UI 清晰度"],
            ["Android 真机流畅度", "Pending", "新包在代表性手机上记录帧时间后关闭"],
            ["Penpot 页面更新", "N/A", "本任务锁定 UE，不改页面结构"],
        ],
        [2.2, 1.55, 2.85],
    )

    doc.core_properties.title = "渲染清晰度与战斗性能优化"
    doc.core_properties.subject = "移动端竖屏有效分辨率、原生像素文字与 72 单位战斗性能技术合同"
    doc.core_properties.author = "Codex Game Studio"
    doc.core_properties.keywords = "Godot, Android, UI clarity, native pixels, battle performance"
    return doc


def main() -> int:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
