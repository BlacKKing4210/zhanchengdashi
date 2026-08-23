from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_OVERRIDE = os.environ.get("ZC_FEATURE_TEMPLATE", "").strip()
TEMPLATE = (
    Path(TEMPLATE_OVERRIDE)
    if TEMPLATE_OVERRIDE
    else Path.home() / ".codex" / "skills" / "game-feature-design-docs" / "assets" / "general-feature-design-template.docx"
)
OUTPUT = ROOT / "docs" / "PLAYER_ACCOUNT_IDENTITY_AVATAR_DESIGN_v2.0.docx"
HUMAN_BOARD = ROOT / "output" / "visual_concepts" / "account-avatar-v2" / "ACCOUNT_AVATAR_HUMAN_OPTIONS_BOARD_v1.png"

INK = "172033"
BLUE = "2E74B5"
TEAL = "197A78"
GRAY = "5B6574"
RED = "C23B3B"
GREEN = "2F8F5B"
ORANGE = "D97706"
LIGHT_BLUE = "EAF3FA"
LIGHT_GOLD = "FFF4D6"
LIGHT_GRAY = "F2F4F7"
LIGHT_RED = "FDECEC"
LIGHT_GREEN = "EAF6EF"
WHITE = "FFFFFF"


def set_run_font(run, size: float = 10.3, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Microsoft YaHei"
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 9.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.02
    set_run_font(paragraph.add_run(str(text)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def keep_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=9.1)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    repeat_header(table.rows[0])
    keep_row(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
        keep_row(table.rows[-1])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_body(doc: Document, text: str, *, bold: bool = False, color: str = INK) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.13
    set_run_font(paragraph.add_run(text), bold=bold, color=color)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text))


def add_step(doc: Document, index: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(f"{index}.  {text}"))


def add_callout(doc: Document, text: str, *, fill: str = LIGHT_GOLD, text_color: str = INK) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_text(table.cell(0, 0), text, bold=True, color=text_color, size=10.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "在 Word 中右键更新目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.3)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 13, 7),
        ("Heading 2", 12.5, TEAL, 9, 5),
        ("Heading 3", 11, GRAY, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.3)

    header = section.header.paragraphs[0]
    header.text = "战城大师｜玩家账号、用户名与头像系统"
    set_run_font(header.runs[0], 8.3, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "F-ZC-AUTH-001｜v2.0｜2026-08-23"
    set_run_font(footer.runs[0], 8.3, color=GRAY)


def add_title_page(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(32)
    title.paragraph_format.space_after = Pt(7)
    set_run_font(title.add_run("玩家账号、用户名与头像系统"), 24, True, INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("简单注册登录｜稳定账号 ID｜可修改用户名｜预设头像｜凭据复制"), 12.2, False, BLUE)

    add_table(
        doc,
        ["字段", "内容"],
        [
            ["项目 / 功能", "战城大师 / F-ZC-AUTH-001"],
            ["文档版本", "v2.0（本地实现基线）"],
            ["负责人", "唯一制作人：项目制作人；执行：codex-primary"],
            ["实现状态", "本地实现授权；人物头像整板、FigJam 与阿里云生产安全门待验收"],
            ["兼容目标", "保留既有 UserID、进度、密码哈希、设备绑定和账号切换"],
            ["运行时头像", "首版开放 12 个现有动物头像；4 个人物候选保持 NOT_RUNTIME"],
        ],
        [1.6, 5.0],
    )
    add_callout(
        doc,
        "核心原则：账号 ID 用于登录且稳定不变；用户名只用于玩家展示且允许重名；服务器永不保存或返回明文旧密码。",
        fill=LIGHT_BLUE,
    )
    doc.add_page_break()
    doc.add_heading("目录", level=1)
    add_toc(doc)
    doc.add_page_break()


def build_document() -> Document:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"General feature template missing: {TEMPLATE}")
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    style_document(doc)
    add_title_page(doc)

    doc.add_heading("1. 版本与决策状态", level=1)
    add_table(
        doc,
        ["版本", "日期", "状态", "变更"],
        [
            ["v1.5", "2026-08-13", "既有基线", "设备账号、账号密码登录、服务器权威资料、复制凭据"],
            ["v2.0", "2026-08-23", "本地实现", "新增独立用户名、预设头像、身份 revision、旧记录兼容迁移"],
        ],
        [0.8, 1.1, 1.1, 3.6],
    )
    add_callout(
        doc,
        "PENDING：可编辑 FigJam UE 源当前因连接器不可用未创建；人物头像整板须制作人书面批准后才能切图和接入；本任务无阿里云远程写授权。",
        fill=LIGHT_RED,
        text_color=RED,
    )

    doc.add_heading("2. 目标、非目标与成功标准", level=1)
    doc.add_heading("2.1 目标", level=2)
    for item in (
        "玩家第一次打开游戏即可自动获得可登录账号，不要求先填写复杂表单；只选择用户名和头像即可完成身份设置。",
        "账号 ID、UserID、用户名三者职责清晰：登录、内部主键、玩家展示互不混用。",
        "玩家可直接使用账号 ID + 密码登录其他设备，并可复制账号、复制密码或复制全部。",
        "房间、账号中心和账号切换列表统一优先显示用户名，必要位置附短 UserID 识别。",
        "既有账号无损迁移；服务器对用户名和头像实施白名单、长度、控制字符与 revision 校验。",
    ):
        add_bullet(doc, item)
    doc.add_heading("2.2 非目标", level=2)
    for item in (
        "不提供自由上传头像、照片审核或完整捏脸系统。",
        "不把用户名当作登录标识，不要求用户名全局唯一。",
        "不在客户端或服务器恢复历史明文密码；只有本次创建或重新验证后的前台会话可复制密码。",
        "本版本不部署阿里云、不迁移线上账号、不宣称生产级口令安全门已通过。",
        "不复制《模拟人生》的角色、服装、UI、图标或素材，只借鉴预设身份表达与低门槛个性化方法。",
    ):
        add_bullet(doc, item)
    doc.add_heading("2.3 可量化成功标准", level=2)
    add_table(
        doc,
        ["指标", "验收阈值"],
        [
            ["首次身份设置", "用户名 + 头像两项，一次确认完成；不再要求另填账号 ID"],
            ["登录路径", "账号 ID + 密码，两项输入；用户名不能登录"],
            ["兼容迁移", "旧记录的 UserID、资料、密码派生值、安装绑定逐字段不变"],
            ["头像首发", "12 个动物头像全部可选且资源存在；人物 4 候选均不进入运行时"],
            ["身份并发", "expected_revision 不匹配时拒绝覆盖并返回最新身份"],
            ["玩家可见证据", "账号中心截图显示独立用户名、账号 ID、UserID、头像和三个复制入口"],
        ],
        [1.5, 5.1],
    )

    doc.add_heading("3. 竞品与安全研究结论", level=1)
    add_table(
        doc,
        ["来源", "可借鉴原则", "本项目落地"],
        [
            ["PlayFab 账号关联", "先低摩擦进入，再关联可恢复身份", "设备账号自动创建；投入后可用账号密码跨设备登录"],
            ["Supercell ID", "少步骤、跨设备、支持多账号切换", "保留本机档案集合和直接切换"],
            ["Steam / Roblox / Nintendo", "稳定账号身份与可变展示身份分离", "account 与 user_id 稳定；username 可修改、允许重名"],
            ["The Sims Create-a-Sim", "用预设快速表达身份，再逐步调整", "有限动物/人物头像预设；不做复杂捏脸"],
            ["OWASP", "TLS、限速、现代密码 KDF、统一失败信息", "生产门禁列为必做；现有 12000 轮 SHA-256 仅兼容旧记录"],
        ],
        [1.5, 2.4, 2.7],
    )
    add_body(doc, "外部资料仅用于原则研究，最终字段、流程、UI、命名与资产均为本项目原创实现。", color=GRAY)

    doc.add_heading("4. 身份模型与权威数据", level=1)
    add_table(
        doc,
        ["字段", "职责", "规则", "权威 / 可见性"],
        [
            ["user_id", "服务器内部主键", "创建后不可修改；不作为玩家输入", "服务器；必要位置显示短尾号"],
            ["account", "登录账号 ID", "唯一、稳定、可复制；不能用作房间显示名", "服务器 + 当前客户端凭据"],
            ["username", "玩家展示名", "2–16 个字符；去首尾空白；拒绝控制字符；允许重名", "服务器；房间/UI 主显示"],
            ["avatar_id", "预设头像 ID", "仅服务器运行时白名单可用", "服务器；客户端按目录映射资源"],
            ["identity_revision", "身份并发版本", "从 1 递增；更新必须提交 expected_revision", "服务器"],
            ["identity_complete", "首次身份设置状态", "自动设备账号初始为 false；首次确认后为 true；旧账号迁移为 true", "服务器；驱动首次完成按钮与提示"],
            ["password_plain", "本次输入或生成口令", "只在当前前台会话短暂存在；断线/切换/失焦清除", "客户端内存；禁止持久化和回传"],
            ["password_hash/salt", "登录校验", "现有记录兼容；生产前迁移到版本化现代 KDF", "服务器；永不返回客户端"],
        ],
        [1.1, 1.4, 2.6, 1.5],
    )
    doc.add_heading("4.1 首版动物头像目录", level=2)
    add_table(
        doc,
        ["avatar_id", "显示名", "既有资源"],
        [
            ["animal_cat", "猫", "assets/card_art/animals/cat.png"],
            ["animal_dog", "狗", "assets/card_art/animals/dog.png"],
            ["animal_fox", "狐狸", "assets/card_art/animals/fox.png"],
            ["animal_rabbit", "兔子", "assets/card_art/animals/rabbit.png"],
            ["animal_tiger", "老虎", "assets/card_art/animals/tiger.png"],
            ["animal_lion", "狮子", "assets/card_art/animals/lion.png"],
            ["animal_elephant", "大象", "assets/card_art/animals/elephant.png"],
            ["animal_penguin", "企鹅", "assets/card_art/animals/penguin.png"],
            ["animal_otter", "水獭", "assets/card_art/animals/otter.png"],
            ["animal_squirrel", "松鼠", "assets/card_art/animals/squirrel.png"],
            ["animal_hamster", "仓鼠", "assets/card_art/animals/hamster.png"],
            ["animal_monkey", "猴子", "assets/card_art/animals/monkey.png"],
        ],
        [1.4, 1.0, 4.2],
    )

    doc.add_heading("5. 核心玩家流程", level=1)
    doc.add_heading("5.1 首次打开 / 自动注册", level=2)
    for index, text in enumerate(
        (
            "客户端生成安装 ID，并向服务器请求设备账号。",
            "服务器生成不可变 UserID、唯一账号 ID 和随机密码；只在首次响应返回一次账号 ID 与密码。",
            "客户端自动进入身份设置：默认动物头像已选中，用户名输入框预填“新玩家 + UserID 后四位”。",
            "玩家修改用户名或头像后点击“完成注册”；客户端提交 session_token、username、avatar_id、expected_revision。",
            "服务器验证并保存身份；账号中心显示头像、用户名、账号 ID、UserID 与复制按钮。",
        ),
        1,
    ):
        add_step(doc, index, text)
    doc.add_heading("5.2 已有玩家登录", level=2)
    for index, text in enumerate(
        (
            "玩家输入账号 ID 和密码；用户名不能作为登录输入。",
            "服务器统一校验并签发短会话，返回 user_id、account、username、avatar_id、identity_revision 与玩家资料。",
            "客户端把该账号加入本机可切换档案并设为当前账号；本次输入密码仅保留在前台内存。",
            "失败时使用统一错误提示，不暴露账号是否存在；达到限速门槛后暂缓重试。",
        ),
        1,
    ):
        add_step(doc, index, text)
    doc.add_heading("5.3 凭据复制", level=2)
    add_table(
        doc,
        ["入口", "可用条件", "写入剪贴板"],
        [
            ["复制账号", "当前会话已认证", "账号：<account>"],
            ["复制密码", "本次创建或重新登录后仍持有前台明文", "密码：<current plaintext>"],
            ["复制全部", "同时满足账号与密码条件", "账号：... 换行 密码：..."],
        ],
        [1.3, 3.2, 2.1],
    )
    add_body(doc, "若当前会话没有明文密码，“复制密码/复制全部”不读取服务器或磁盘，而是提示“请重新登录验证后复制”。", bold=True)

    doc.add_heading("6. UI/UE 规格", level=1)
    add_callout(doc, "最终可编辑 UE 源：FigJam（PENDING_CONNECTOR_UNAVAILABLE）。本节是实现合同，不冒充最终 FigJam 设计源。", fill=LIGHT_RED, text_color=RED)
    doc.add_heading("6.1 账号中心信息层级", level=2)
    add_table(
        doc,
        ["优先级", "区域", "内容 / 行为"],
        [
            ["P0", "身份头部", "圆形头像 + 大号用户名；不把账号 ID 当名字"],
            ["P0", "登录信息", "账号 ID、UserID 分两行；各自可复制但不抢主视觉"],
            ["P0", "主操作", "未完成身份时“完成注册”；已完成时“保存资料”"],
            ["P1", "头像选择", "横向/网格显示 12 个动物头像；当前选中态清晰"],
            ["P1", "凭据工具", "复制账号、复制密码、复制全部三个独立按钮"],
            ["P2", "账号切换", "列表显示头像、用户名、账号 ID 尾号；切换不丢资料"],
        ],
        [0.7, 1.5, 4.4],
    )
    doc.add_heading("6.2 页面状态矩阵", level=2)
    add_table(
        doc,
        ["状态", "用户名/头像", "主按钮", "复制密码", "提示"],
        [
            ["首次自动账号", "可编辑；默认值", "完成注册", "可用", "先保存账号信息"],
            ["已登录且有前台密码", "可编辑", "保存资料", "可用", "密码只在本次会话可复制"],
            ["自动恢复登录", "可编辑", "保存资料", "不可用", "重新登录验证后可复制密码"],
            ["保存中", "锁定", "保存中…", "保持原状态", "防重复提交"],
            ["revision 冲突", "刷新为服务器最新值", "重新编辑", "保持原状态", "资料已在其他设备更新"],
            ["未认证/断线", "只读或隐藏", "重新连接", "不可用并清空", "连接服务器后操作"],
        ],
        [1.4, 1.5, 1.2, 1.1, 1.4],
    )
    doc.add_heading("6.3 输入与错误文案", level=2)
    add_table(
        doc,
        ["条件", "玩家提示"],
        [
            ["用户名少于 2 个字符", "用户名至少 2 个字符"],
            ["用户名超过 16 个字符", "用户名最多 16 个字符"],
            ["用户名含控制字符", "用户名包含不可用字符，请修改"],
            ["头像不在运行时白名单", "该头像暂未开放"],
            ["身份 revision 冲突", "资料已在其他设备更新，已载入最新资料"],
            ["登录失败", "账号或密码错误，请重试"],
            ["无前台密码", "请重新登录验证后复制密码"],
        ],
        [2.4, 4.2],
    )

    doc.add_heading("7. 人物头像整板资产合同", level=1)
    add_table(
        doc,
        ["项目", "要求"],
        [
            ["整板", "一张高分辨率 2×2 评审板；4 个原创通用卡通人物头像；构图、光照和线条一致"],
            ["角色差异", "不同脸型、肤色、发型与服装色；不绑定职业或战斗强度；避免性别刻板化"],
            ["风格", "简洁高质量 2D 游戏头像；友好、可读、色块克制；只借鉴预设身份表达的方法"],
            ["禁止", "不得出现《模拟人生》角色、绿色晶锥标识、UI、Logo、服装复刻或可识别素材"],
            ["状态", "整板和 4 个候选单元均为 NOT_RUNTIME；制作人书面批准前不切图、不去底、不导入 Godot"],
            ["文件 / SHA-256", "output/visual_concepts/account-avatar-v2/ACCOUNT_AVATAR_HUMAN_OPTIONS_BOARD_v1.png / 4C5828588574E2842A34A7F47C9A8B181F115C9F1D60656D8B8DDF1CF1DDBC45"],
            ["批准后流程", "冻结整板 SHA-256 → 切图 → 透明度/边缘/尺寸 QA → avatar_id 映射 → 运行时截图验收"],
        ],
        [1.4, 5.2],
    )
    if HUMAN_BOARD.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(HUMAN_BOARD), width=Inches(6.2))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(caption.add_run("人物头像候选整板 v1｜NOT_RUNTIME｜待制作人书面批准"), 9.0, True, RED)
    else:
        add_callout(doc, f"整板待生成：{HUMAN_BOARD.relative_to(ROOT).as_posix()}｜NOT_RUNTIME", fill=LIGHT_GOLD)

    doc.add_heading("8. 服务器、RPC 与迁移合同", level=1)
    doc.add_heading("8.1 服务器记录迁移", level=2)
    add_table(
        doc,
        ["旧记录条件", "迁移结果", "禁止变化"],
        [
            ["缺 username", "优先使用 account；否则“新玩家 + UserID 后四位”", "user_id、profile、hash/salt、installations"],
            ["缺 avatar_id", "写入 animal_cat", "进度与登录凭据"],
            ["缺 identity_revision", "写入 1", "已有 revision 不回退"],
            ["缺 identity_complete", "旧账号写入 true；新设备账号在首次确认前保持 false", "不得要求旧玩家重复完成首次设置"],
            ["字段无效", "按同一默认规则修复并持久化", "不得删除账号或重置资料"],
        ],
        [1.7, 2.7, 2.2],
    )
    doc.add_heading("8.2 RPC/返回字段", level=2)
    add_table(
        doc,
        ["操作", "请求", "成功返回", "拒绝条件"],
        [
            ["login", "account, password, install_id", "session + identity + profile", "统一认证失败 / 限速"],
            ["update_identity", "session, username, avatar_id, expected_revision", "identity + new revision", "会话无效 / 字段无效 / revision 冲突"],
            ["profile_for_session", "session", "profile + identity", "会话无效"],
            ["account_summaries", "install credential", "avatar, username, account, user_id, profile summary", "不属于安装集合"],
        ],
        [1.3, 2.0, 1.9, 1.5],
    )
    add_body(doc, "服务器响应永不包含 salt、password_hash、旧明文密码或其他账号的凭据。客户端提交 username 时不能连带覆盖玩家 profile。", bold=True)

    doc.add_heading("9. 安全、隐私与发布门禁", level=1)
    for item in (
        "现有 12000 轮 SHA-256 只用于兼容旧记录；生产前必须引入 password_kdf_version 与现代 KDF（优先 Argon2id，参数按服务器能力标定）。",
        "公网注册、登录、改密和恢复必须走校验证书的加密认证通道；普通明文 UDP 认证不得通过发行门。",
        "登录、注册、身份修改和恢复接口必须做账号/IP/设备维度限速与审计；错误信息不得枚举账号。",
        "剪贴板只在玩家主动操作后写入；客户端 60 秒后且内容未被改写时尝试清空。公共设备需提示风险。",
        "本任务无阿里云远程写授权，不读取或修改 staging 的真实玩家行；线上迁移需独立备份、回滚和外部客户端验收。",
    ):
        add_bullet(doc, item)
    add_table(
        doc,
        ["门禁", "当前状态", "通过条件"],
        [
            ["本地功能", "本版本执行", "代码、单测、项目解析、玩家可见截图"],
            ["人物头像", "PENDING", "制作人批准整板，之后切图和运行时 QA"],
            ["FigJam UE", "PENDING", "连接器恢复后创建可编辑流程和页面规格"],
            ["阿里云 staging", "NOT AUTHORIZED", "明确环境/变更授权 + 备份回滚 + 真实外网测试"],
            ["生产安全", "BLOCKED", "现代 KDF、TLS、限速、恢复、撤销、渗透与抓包验收"],
        ],
        [1.4, 1.5, 3.7],
    )

    doc.add_heading("10. QA 验收用例", level=1)
    add_table(
        doc,
        ["编号", "用例", "预期"],
        [
            ["AUTH2-001", "载入 v1 旧账号记录", "自动补 username/avatar/revision；其余权威字段字节级不变"],
            ["AUTH2-002", "两个账号保存相同用户名", "均成功；UserID 和 account 仍各自独立"],
            ["AUTH2-003", "用 username + 密码登录", "拒绝；只有 account + 密码可登录"],
            ["AUTH2-004", "提交 1/17 字符、控制字符用户名", "拒绝并返回稳定错误码"],
            ["AUTH2-005", "提交未批准人物 avatar_id", "服务器拒绝，不写入记录"],
            ["AUTH2-006", "两个客户端用同一 revision 更新", "首个成功；后一个冲突并拿到最新身份"],
            ["AUTH2-007", "自动登录后复制密码", "不可用；提示重新登录，不读取磁盘或服务器"],
            ["AUTH2-008", "重新登录后复制账号/密码/全部", "内容准确；失焦/切换/断线清除明文"],
            ["AUTH2-009", "房间与账号切换列表", "显示 username + avatar，不再使用“玩家X”或 account 当名字"],
            ["AUTH2-010", "GDScript 与项目解析", "tab 缩进检查、目标测试、全项目 headless 解析通过"],
        ],
        [1.0, 2.8, 2.8],
    )

    doc.add_heading("11. 数据、埋点与运营边界", level=1)
    add_table(
        doc,
        ["事件", "字段（不含敏感值）", "用途"],
        [
            ["identity_setup_complete", "avatar_category, name_length, elapsed_ms", "评估首次设置完成率"],
            ["identity_update_result", "result_code, avatar_changed, name_changed", "发现校验/冲突问题"],
            ["credential_copy", "copy_type, has_foreground_password", "评估复制功能需求；绝不记录内容"],
            ["manual_login_result", "result_code, latency_bucket", "监控成功率与性能；不记录账号/密码"],
        ],
        [1.8, 3.0, 1.8],
    )
    add_body(doc, "用户名属于玩家可见内容，后台搜索与审核应按项目隐私规则处理；日志和埋点不得写入账号、密码、token、install_id 或完整 UserID。", bold=True)

    doc.add_heading("12. 外部参考与原始来源", level=1)
    sources = (
        ("Microsoft PlayFab｜Link multiple authentication providers", "https://learn.microsoft.com/en-us/gaming/playfab/identity/player-identity/platform-specific-authentication/linking-accounts"),
        ("Supercell｜Supercell ID", "https://supercell.com/en/supercell-id/"),
        ("Roblox Support｜Changing Your Display Name", "https://en.help.roblox.com/hc/en-us/articles/4401938870292-Changing-Your-Display-Name"),
        ("Nintendo Support｜Nintendo Account nickname/icon", "https://en-americas-support.nintendo.com/app/answers/detail/a_id/15988"),
        ("EA｜The Sims 4 Create-a-Sim", "https://www.ea.com/games/the-sims/the-sims-4/news/create-a-sim-demo"),
        ("OWASP｜Authentication Cheat Sheet", "https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html"),
        ("OWASP｜Password Storage Cheat Sheet", "https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html"),
    )
    for title, url in sources:
        paragraph = doc.add_paragraph(style="List Bullet")
        set_run_font(paragraph.add_run(f"{title}："), bold=True)
        run = paragraph.add_run(url)
        set_run_font(run, color=BLUE)

    doc.add_heading("13. 交付与待确认清单", level=1)
    add_table(
        doc,
        ["项目", "本版本动作", "交付判断"],
        [
            ["设计与规则", "本 DOCX 固化", "可评审"],
            ["动物头像", "复用既有 12 张运行时资源", "纳入本地实现"],
            ["人物头像", "只交整板", "NOT_RUNTIME，等待制作人选择"],
            ["客户端/服务器", "实现 username/avatar/revision 与登录/复制/UI", "以测试和截图为准"],
            ["FigJam", "当前连接器不可用", "PENDING，不冒充完成"],
            ["阿里云", "不写入", "另行授权后处理"],
        ],
        [1.5, 3.0, 2.1],
    )
    doc.core_properties.title = "玩家账号、用户名与头像系统设计"
    doc.core_properties.subject = "账号 ID 与用户名分离、预设头像、简单注册登录与凭据复制"
    doc.core_properties.author = "Codex Game Studio"
    doc.core_properties.comments = "Based on the approved general game feature design template."
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
