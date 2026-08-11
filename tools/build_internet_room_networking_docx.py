#!/usr/bin/env python3
"""Build the editable Word specification for internet room networking."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_game_audio_design_docx as base


ROOT = Path(__file__).resolve().parents[1]


def format_code_blocks(doc: Document) -> None:
    in_code_block = False
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("```") or text in {"`", "``"}:
            in_code_block = not in_code_block
            element = paragraph._element
            element.getparent().remove(element)
            continue
        if not in_code_block:
            continue
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.08)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        p_pr = paragraph._p.get_or_add_pPr()
        shading = p_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            p_pr.append(shading)
        shading.set(qn("w:fill"), base.LIGHT_GRAY)
        for run in paragraph.runs:
            base.set_run_font(run, size=8.5, color=base.DARK_BLUE)


def add_masthead(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("丛林法则  |  公网多人房间技术规格")
    base.set_run_font(run, size=9, bold=True, color=base.MUTED)
    base.add_page_number(section.footer.paragraphs[0])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("公网房间联网规格")
    base.set_run_font(run, size=23, bold=True, color=base.DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    base.add_inline_runs(
        subtitle,
        "ENet 公网接入、权威房间槽位、准备开局与战斗快照路由",
        size=12,
        color=base.MUTED,
    )
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    base.add_inline_runs(meta, "版本 v1.1  |  状态：首版实现基线  |  2026-07-13", size=10, color=base.MUTED)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), base.BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def main() -> None:
    base.SOURCE = ROOT / "docs" / "INTERNET_ROOM_NETWORKING.md"
    base.OUTPUT = ROOT / "docs" / "INTERNET_ROOM_NETWORKING.docx"
    base.add_masthead = add_masthead
    base.build_document()
    doc = Document(base.OUTPUT)
    format_code_blocks(doc)
    properties = doc.core_properties
    properties.title = "公网房间联网规格"
    properties.subject = "丛林法则 ENet 公网房间、槽位与战斗同步"
    properties.author = "Codex Game Studio"
    properties.keywords = "Godot, ENet, 公网房间, 1V1, 2V2, 3V3, 战斗同步"
    doc.save(base.OUTPUT)
    print(base.OUTPUT)


if __name__ == "__main__":
    main()
